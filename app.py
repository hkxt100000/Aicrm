"""
企业微信 CRM 后端服务
"""
import os
import time
import sqlite3
from pathlib import Path
from typing import List, Dict, Optional, Any
from datetime import datetime

from fastapi import FastAPI, HTTPException, Depends, Query
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from config import DB_PATH, PORT, HOST, API_TOKEN
from wecom_client import wecom_client
from exporter import CustomerExporter
from sync_service import SyncService
from group_tags_api import router as group_tags_router

# 创建同步服务实例（10线程并发）
sync_service = SyncService(wecom_client, max_workers=10)

# 创建 FastAPI 应用
app = FastAPI(title="企业微信 CRM", version="1.0")

# CORS 配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 引入客户群标签路由
app.include_router(group_tags_router)

# 引入企微机器人路由
from bot_api import router as bot_router
app.include_router(bot_router, prefix="/api/bot", tags=["企微机器人"])


# ========== 数据库初始化 ==========

def init_database():
    """初始化数据库"""
    Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    # 创建客户表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS customers (
            id TEXT PRIMARY KEY,
            name TEXT,
            avatar TEXT,
            gender INTEGER DEFAULT 0,
            type INTEGER DEFAULT 1,
            unionid TEXT,
            position TEXT,
            corp_name TEXT,
            corp_full_name TEXT,
            owner_userid TEXT,
            owner_name TEXT,
            add_time INTEGER,
            tags TEXT,
            remark TEXT,
            follow_status TEXT DEFAULT '未跟进',
            satisfaction TEXT DEFAULT '未评价',
            stage TEXT DEFAULT '潜在客户',
            source TEXT,
            groups TEXT,
            last_contact_time INTEGER,
            created_at INTEGER DEFAULT 0,
            updated_at INTEGER DEFAULT 0
        )
    """)
    
    # 创建员工表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS employees (
            id TEXT PRIMARY KEY,
            name TEXT,
            avatar TEXT,
            mobile TEXT,
            email TEXT,
            department TEXT,
            position TEXT,
            status INTEGER DEFAULT 1,
            customer_count INTEGER DEFAULT 0,
            created_at INTEGER DEFAULT 0,
            updated_at INTEGER DEFAULT 0
        )
    """)
    
    # 创建标签表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS customer_tags (
            id TEXT PRIMARY KEY,
            name TEXT,
            group_name TEXT,
            order_num INTEGER DEFAULT 0,
            color TEXT,
            created_at INTEGER DEFAULT 0
        )
    """)
    
    # 创建跟进记录表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS follow_records (
            id TEXT PRIMARY KEY,
            customer_id TEXT,
            employee_id TEXT,
            follow_time INTEGER,
            follow_type TEXT,
            content TEXT,
            next_follow_time INTEGER,
            created_at INTEGER DEFAULT 0
        )
    """)
    
    # 创建配置表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS config (
            key TEXT PRIMARY KEY,
            value TEXT,
            updated_at INTEGER DEFAULT 0
        )
    """)
    
    # 创建智能表格主表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS smart_spreadsheets (
            id TEXT PRIMARY KEY,
            docid TEXT NOT NULL,
            sheet_id TEXT,
            name TEXT NOT NULL,
            data_type TEXT DEFAULT 'order',
            data_scope TEXT DEFAULT 'global',
            supplier_code TEXT,
            file_name TEXT,
            file_path TEXT,
            fields_config TEXT,
            field_mapping TEXT,
            sync_config TEXT,
            row_count INTEGER DEFAULT 0,
            col_count INTEGER DEFAULT 0,
            created_at INTEGER DEFAULT 0,
            updated_at INTEGER DEFAULT 0,
            last_sync_at INTEGER DEFAULT 0,
            version INTEGER DEFAULT 1,
            data_hash TEXT,
            status TEXT DEFAULT 'active',
            url TEXT
        )
    """)
    
    # 创建表格数据表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS spreadsheet_data (
            id TEXT PRIMARY KEY,
            spreadsheet_id TEXT NOT NULL,
            row_index INTEGER NOT NULL,
            col_index INTEGER NOT NULL,
            col_name TEXT,
            value TEXT,
            version INTEGER DEFAULT 1,
            created_at INTEGER DEFAULT 0,
            updated_at INTEGER DEFAULT 0,
            FOREIGN KEY (spreadsheet_id) REFERENCES smart_spreadsheets(id)
        )
    """)
    
    # 创建表格字段模板表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS field_templates (
            id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            data_type TEXT NOT NULL,
            fields_config TEXT NOT NULL,
            description TEXT,
            is_system INTEGER DEFAULT 0,
            created_at INTEGER DEFAULT 0,
            updated_at INTEGER DEFAULT 0
        )
    """)
    
    # 创建同步日志表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS sync_logs (
            id TEXT PRIMARY KEY,
            spreadsheet_id TEXT NOT NULL,
            sync_type TEXT,
            sync_direction TEXT,
            changes_count INTEGER DEFAULT 0,
            status TEXT,
            error_message TEXT,
            sync_data TEXT,
            created_at INTEGER DEFAULT 0,
            FOREIGN KEY (spreadsheet_id) REFERENCES smart_spreadsheets(id)
        )
    """)
    
    # 创建客户群表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS customer_groups (
            chat_id TEXT PRIMARY KEY,
            name TEXT,
            owner_userid TEXT,
            owner_name TEXT,
            notice TEXT,
            member_count INTEGER DEFAULT 0,
            external_member_count INTEGER DEFAULT 0,
            internal_member_count INTEGER DEFAULT 0,
            admin_list TEXT,
            group_type TEXT DEFAULT 'external',
            status INTEGER DEFAULT 0,
            version INTEGER DEFAULT 0,
            create_time INTEGER DEFAULT 0,
            last_sync_time INTEGER DEFAULT 0,
            created_at INTEGER DEFAULT 0,
            updated_at INTEGER DEFAULT 0
        )
    """)
    
    conn.commit()
    
    # 初始化系统字段模板
    import json
    cursor = conn.cursor()
    
    # 检查是否已有模板
    cursor.execute("SELECT COUNT(*) FROM field_templates WHERE is_system = 1")
    count = cursor.fetchone()[0]
    
    if count == 0:
        print("[数据库] 初始化字段模板...")
        
        # 订单完整模板（43字段）- 严格按照用户提供的字段列表
        order_full_fields = [
            {"wecom_name": "下单时间", "system_name": "order_time", "type": "datetime", "editable": False},
            {"wecom_name": "更新时间", "system_name": "update_time", "type": "datetime", "editable": False},
            {"wecom_name": "订单ID", "system_name": "order_id", "type": "text", "editable": False},
            {"wecom_name": "下游订单号", "system_name": "downstream_order_no", "type": "text", "editable": False},
            {"wecom_name": "上游订单号", "system_name": "upstream_order_no", "type": "text", "editable": False},
            {"wecom_name": "订单来源", "system_name": "order_source", "type": "text", "editable": False},
            {"wecom_name": "订单归属人", "system_name": "order_owner", "type": "text", "editable": False},
            {"wecom_name": "订单归属上级", "system_name": "order_owner_superior", "type": "text", "editable": False},
            {"wecom_name": "商务账号", "system_name": "business_account", "type": "text", "editable": False},
            {"wecom_name": "供应商", "system_name": "supplier_name", "type": "text", "editable": False},
            {"wecom_name": "供应商对接群", "system_name": "supplier_contact_group", "type": "text", "editable": False},
            {"wecom_name": "订单状态", "system_name": "order_status", "type": "text", "editable": False},
            {"wecom_name": "商品编码", "system_name": "product_code", "type": "text", "editable": False},
            {"wecom_name": "商品名称", "system_name": "product_name", "type": "text", "editable": False},
            {"wecom_name": "证件姓名", "system_name": "id_name", "type": "text", "editable": False},
            {"wecom_name": "用户微信", "system_name": "user_wechat", "type": "text", "editable": False},
            {"wecom_name": "年龄", "system_name": "age", "type": "number", "editable": False},
            {"wecom_name": "性别", "system_name": "gender", "type": "text", "editable": False},
            {"wecom_name": "收件号码", "system_name": "receiver_phone", "type": "text", "editable": False},
            {"wecom_name": "省", "system_name": "province", "type": "text", "editable": False},
            {"wecom_name": "市", "system_name": "city", "type": "text", "editable": False},
            {"wecom_name": "区", "system_name": "district", "type": "text", "editable": False},
            {"wecom_name": "详情描述", "system_name": "detail_description", "type": "text", "editable": False},
            {"wecom_name": "生产号码", "system_name": "production_number", "type": "text", "editable": False},
            {"wecom_name": "快递公司", "system_name": "express_company", "type": "text", "editable": False},
            {"wecom_name": "物流单号", "system_name": "express_no", "type": "text", "editable": False},
            {"wecom_name": "预存话费", "system_name": "prepaid_balance", "type": "number", "editable": False},
            {"wecom_name": "交易状态", "system_name": "transaction_status", "type": "text", "editable": False},
            {"wecom_name": "交易单号", "system_name": "transaction_no", "type": "text", "editable": False},
            {"wecom_name": "激活状态", "system_name": "activation_status", "type": "text", "editable": False},
            {"wecom_name": "激活时间", "system_name": "activation_time", "type": "datetime", "editable": False},
            {"wecom_name": "首充状态", "system_name": "first_recharge_status", "type": "text", "editable": False},
            {"wecom_name": "首充金额", "system_name": "first_recharge_amount", "type": "number", "editable": False},
            {"wecom_name": "分佣状态", "system_name": "commission_status", "type": "text", "editable": False},
            {"wecom_name": "平台服务费", "system_name": "platform_service_fee", "type": "number", "editable": False},
            {"wecom_name": "上级居间费", "system_name": "superior_intermediary_fee", "type": "number", "editable": False},
            {"wecom_name": "店铺佣金", "system_name": "shop_commission", "type": "number", "editable": False},
            {"wecom_name": "推客赏金", "system_name": "promoter_reward", "type": "number", "editable": False},
            {"wecom_name": "顾客奖金", "system_name": "customer_reward", "type": "number", "editable": False},
            {"wecom_name": "结算模式", "system_name": "settlement_mode", "type": "text", "editable": False},
            {"wecom_name": "结算规则", "system_name": "settlement_rule", "type": "text", "editable": False},
            {"wecom_name": "结算时间", "system_name": "settlement_time", "type": "datetime", "editable": False},
            {"wecom_name": "结算说明", "system_name": "settlement_description", "type": "text", "editable": False}
        ]
        
        # 订单精简模板（10字段）
        order_simple_fields = [
            {"wecom_name": "下单时间", "system_name": "order_time", "type": "datetime", "editable": False},
            {"wecom_name": "订单ID", "system_name": "order_id", "type": "text", "editable": False},
            {"wecom_name": "客户姓名", "system_name": "customer_name", "type": "text", "editable": False},
            {"wecom_name": "客户电话", "system_name": "customer_phone", "type": "text", "editable": False},
            {"wecom_name": "订单金额", "system_name": "order_amount", "type": "number", "editable": False},
            {"wecom_name": "订单状态", "system_name": "order_status", "type": "text", "editable": False},
            {"wecom_name": "支付状态", "system_name": "payment_status", "type": "text", "editable": False},
            {"wecom_name": "发货状态", "system_name": "delivery_status", "type": "text", "editable": False},
            {"wecom_name": "代理商名称", "system_name": "supplier_name", "type": "text", "editable": False},
            {"wecom_name": "订单备注", "system_name": "order_remark", "type": "text", "editable": False}
        ]
        
        # 代理商模板（20字段）
        supplier_fields = [
            {"wecom_name": "代理商编号", "system_name": "supplier_code", "type": "text", "editable": False},
            {"wecom_name": "代理商名称", "system_name": "supplier_name", "type": "text", "editable": False},
            {"wecom_name": "联系人", "system_name": "contact_name", "type": "text", "editable": False},
            {"wecom_name": "联系电话", "system_name": "contact_phone", "type": "text", "editable": False},
            {"wecom_name": "联系邮箱", "system_name": "contact_email", "type": "text", "editable": False},
            {"wecom_name": "公司地址", "system_name": "company_address", "type": "text", "editable": False},
            {"wecom_name": "合作状态", "system_name": "cooperation_status", "type": "text", "editable": False},
            {"wecom_name": "合作等级", "system_name": "cooperation_level", "type": "text", "editable": False},
            {"wecom_name": "合作开始时间", "system_name": "cooperation_start_time", "type": "datetime", "editable": False},
            {"wecom_name": "合作结束时间", "system_name": "cooperation_end_time", "type": "datetime", "editable": False},
            {"wecom_name": "授信额度", "system_name": "credit_limit", "type": "number", "editable": False},
            {"wecom_name": "已用额度", "system_name": "used_credit", "type": "number", "editable": False},
            {"wecom_name": "剩余额度", "system_name": "remaining_credit", "type": "number", "editable": False},
            {"wecom_name": "结算周期", "system_name": "settlement_cycle", "type": "text", "editable": False},
            {"wecom_name": "结算方式", "system_name": "settlement_method", "type": "text", "editable": False},
            {"wecom_name": "累计订单数", "system_name": "total_orders", "type": "number", "editable": False},
            {"wecom_name": "累计金额", "system_name": "total_amount", "type": "number", "editable": False},
            {"wecom_name": "最后下单时间", "system_name": "last_order_time", "type": "datetime", "editable": False},
            {"wecom_name": "客户经理", "system_name": "account_manager", "type": "text", "editable": False},
            {"wecom_name": "备注", "system_name": "remark", "type": "text", "editable": False}
        ]
        
        # 产品模板（15字段）
        product_fields = [
            {"wecom_name": "产品ID", "system_name": "product_id", "type": "text", "editable": False},
            {"wecom_name": "产品名称", "system_name": "product_name", "type": "text", "editable": False},
            {"wecom_name": "产品分类", "system_name": "product_category", "type": "text", "editable": False},
            {"wecom_name": "产品型号", "system_name": "product_model", "type": "text", "editable": False},
            {"wecom_name": "产品规格", "system_name": "product_spec", "type": "text", "editable": False},
            {"wecom_name": "产品单价", "system_name": "product_price", "type": "number", "editable": False},
            {"wecom_name": "成本价", "system_name": "cost_price", "type": "number", "editable": False},
            {"wecom_name": "库存数量", "system_name": "stock_quantity", "type": "number", "editable": False},
            {"wecom_name": "预警库存", "system_name": "warning_stock", "type": "number", "editable": False},
            {"wecom_name": "产品状态", "system_name": "product_status", "type": "text", "editable": False},
            {"wecom_name": "上架时间", "system_name": "online_time", "type": "datetime", "editable": False},
            {"wecom_name": "下架时间", "system_name": "offline_time", "type": "datetime", "editable": False},
            {"wecom_name": "供应商", "system_name": "supplier_name", "type": "text", "editable": False},
            {"wecom_name": "产品描述", "system_name": "product_description", "type": "text", "editable": False},
            {"wecom_name": "备注", "system_name": "remark", "type": "text", "editable": False}
        ]
        
        import time
        now = int(time.time())
        import uuid
        
        templates = [
            (str(uuid.uuid4()), "订单完整模板", "order", json.dumps(order_full_fields, ensure_ascii=False), "包含订单的所有43个字段，适合详细订单管理", 1, now, now),
            (str(uuid.uuid4()), "订单精简模板", "order", json.dumps(order_simple_fields, ensure_ascii=False), "包含订单的10个核心字段，适合快速查看", 1, now, now),
            (str(uuid.uuid4()), "代理商模板", "supplier", json.dumps(supplier_fields, ensure_ascii=False), "代理商信息管理，包含20个字段", 1, now, now),
            (str(uuid.uuid4()), "产品模板", "product", json.dumps(product_fields, ensure_ascii=False), "产品信息管理，包含15个字段", 1, now, now)
        ]
        
        cursor.executemany("""
            INSERT INTO field_templates (id, name, data_type, fields_config, description, is_system, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, templates)
        
        print(f"[数据库] 已初始化 {len(templates)} 个字段模板")
    
    conn.commit()
    conn.close()
    print("[数据库] 初始化完成")

# 启动时初始化数据库
init_database()

# ========== 工具函数 ==========

def get_db():
    """获取数据库连接"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def check_token(token: str = Query(..., alias="api_token")):
    """验证 API Token"""
    if token != API_TOKEN:
        raise HTTPException(status_code=401, detail="Invalid API token")
    return token

def get_wecom_config():
    """获取企业微信配置（从请求头或默认配置）"""
    # 尝试从前端传递的配置获取
    # 如果没有，返回空配置，让 API 使用默认值
    return {
        'corpid': '',
        'contact_secret': '',
        'customer_secret': '',
        'app_secret': '',
        'agentid': ''
    }

# ========== Pydantic 模型 ==========

class SyncCustomersRequest(BaseModel):
    """同步客户请求"""
    force: bool = False  # 是否强制全量同步
    config: Optional[dict] = None  # 企业微信配置

class SyncRequest(BaseModel):
    """通用同步请求"""
    config: Optional[dict] = None  # 企业微信配置

class CustomerFilter(BaseModel):
    """客户筛选"""
    owner_userid: Optional[str] = None
    follow_status: Optional[str] = None
    stage: Optional[str] = None
    satisfaction: Optional[str] = None
    tags: Optional[List[str]] = None
    add_time_start: Optional[int] = None
    add_time_end: Optional[int] = None
    search: Optional[str] = None

# ========== API 路由 ==========

@app.get("/health")
async def health():
    """健康检查"""
    return {
        "status": "ok",
        "version": "1.0",
        "module": "wecom-crm",
        "port": PORT
    }

@app.post("/api/sync/customers")
async def sync_customers(request: SyncCustomersRequest, token: str = Depends(check_token)):
    """
    同步客户数据（增量同步，支持10线程并发和后台队列）
    """
    try:
        print("[API] 开始增量同步客户...")
        
        # 判断同步类型
        task_type = 'full' if request.force else 'incremental'
        
        # 启动后台同步任务
        task_id = sync_service.start_sync_task(
            task_type=task_type,
            config=request.config
        )
        
        print(f"[API] 同步任务已创建: {task_id} (类型: {task_type})")
        
        return {
            "success": True,
            "message": f"同步任务已启动 ({'全量' if task_type == 'full' else '增量'}同步)",
            "task_id": task_id,
            "task_type": task_type
        }
        
    except Exception as e:
        print(f"[API] 启动同步任务失败: {e}")
        return {"success": False, "message": str(e)}

@app.get("/api/sync/status/{task_id}")
async def get_sync_status(task_id: str, token: str = Depends(check_token)):
    """
    获取同步任务状态
    """
    try:
        status = sync_service.get_task_status(task_id)
        
        if not status:
            return {
                "success": False,
                "message": "任务不存在"
            }
        
        return {
            "success": True,
            "data": status
        }
        
    except Exception as e:
        print(f"[API] 获取任务状态失败: {e}")
        return {"success": False, "message": str(e)}

@app.post("/api/sync/stop/{task_id}")
async def stop_sync(task_id: str, token: str = Depends(check_token)):
    """
    停止同步任务
    """
    try:
        result = sync_service.stop_task(task_id)
        
        if result:
            return {
                "success": True,
                "message": "停止信号已发出，任务即将终止"
            }
        else:
            return {
                "success": False,
                "message": "无法停止任务（任务不存在或已结束）"
            }
        
    except Exception as e:
        print(f"[API] 停止任务失败: {e}")
        return {"success": False, "message": str(e)}

@app.get("/api/customers")
async def get_customers(
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    owner_userid: Optional[str] = None,
    user_type: Optional[str] = None,
    add_way: Optional[str] = None,
    date_start: Optional[str] = None,
    date_end: Optional[str] = None,
    wechat_type: Optional[str] = None,
    gender: Optional[str] = None,
    tags: Optional[str] = None,
    provinces: Optional[str] = None,
    search: Optional[str] = None,
    token: str = Depends(check_token)
):
    """
    获取客户列表（分页 + 筛选）
    """
    conn = get_db()
    cursor = conn.cursor()
    
    # 构建查询条件
    where_clauses = []
    params = []
    
    if owner_userid:
        where_clauses.append("owner_userid = ?")
        params.append(owner_userid)
    
    if add_way:
        where_clauses.append("add_way = ?")
        params.append(add_way)
    
    if gender:
        where_clauses.append("gender = ?")
        params.append(gender)
    
    if date_start:
        where_clauses.append("add_time >= ?")
        params.append(int(datetime.strptime(date_start, '%Y-%m-%d').timestamp()))
    
    if date_end:
        where_clauses.append("add_time <= ?")
        params.append(int(datetime.strptime(date_end, '%Y-%m-%d').timestamp()) + 86400)
    
    if search:
        where_clauses.append("(name LIKE ? OR remark LIKE ? OR corp_name LIKE ?)")
        params.extend([f"%{search}%", f"%{search}%", f"%{search}%"])
    
    # 用户属性筛选（通过标签，支持多关键词）
    if user_type:
        if '|' in user_type:
            # 支持多关键词，用 | 分隔，例如 "用户|客户"
            keywords = user_type.split('|')
            user_type_conditions = []
            for keyword in keywords:
                user_type_conditions.append("enterprise_tags LIKE ?")
                params.append(f"%{keyword}%")
            where_clauses.append(f"({' OR '.join(user_type_conditions)})")
        else:
            where_clauses.append("enterprise_tags LIKE ?")
            params.append(f"%{user_type}%")
    
    # 企业标签筛选
    if tags:
        tag_list = tags.split(',')
        tag_conditions = []
        for tag_id in tag_list:
            tag_conditions.append("enterprise_tags LIKE ?")
            params.append(f"%{tag_id}%")
        where_clauses.append(f"({' OR '.join(tag_conditions)})")
    
    # 省份筛选
    if provinces:
        province_list = provinces.split(',')
        province_conditions = []
        for province in province_list:
            province_conditions.append("enterprise_tags LIKE ?")
            params.append(f"%{province}%")
        where_clauses.append(f"({' OR '.join(province_conditions)})")
    
    where_sql = " AND ".join(where_clauses) if where_clauses else "1=1"
    
    # 查询总数
    cursor.execute(f"SELECT COUNT(*) as total FROM customers WHERE {where_sql}", params)
    total = cursor.fetchone()['total']
    
    # 查询列表
    offset = (page - 1) * limit
    cursor.execute(f"""
        SELECT * FROM customers 
        WHERE {where_sql}
        ORDER BY add_time DESC
        LIMIT ? OFFSET ?
    """, params + [limit, offset])
    
    rows = cursor.fetchall()
    customers = [dict(row) for row in rows]
    
    # 解析 JSON 字段
    import json
    for customer in customers:
        try:
            customer['tags'] = json.loads(customer.get('tags') or '[]')
        except:
            customer['tags'] = []
        try:
            customer['groups'] = json.loads(customer.get('groups') or '[]')
        except:
            customer['groups'] = []
        try:
            customer['enterprise_tags'] = json.loads(customer.get('enterprise_tags') or '[]')
        except:
            customer['enterprise_tags'] = []
    
    conn.close()
    
    return {
        "success": True,
        "data": customers,
        "total": total,
        "page": page,
        "limit": limit
    }

@app.get("/api/customers/{customer_id}")
async def get_customer_detail(customer_id: str, token: str = Depends(check_token)):
    """
    获取客户详情
    """
    conn = get_db()
    cursor = conn.cursor()
    
    # 获取客户基本信息
    cursor.execute("SELECT * FROM customers WHERE id = ?", (customer_id,))
    row = cursor.fetchone()
    
    if not row:
        conn.close()
        return {"success": False, "message": "客户不存在"}
    
    customer = dict(row)
    
    # 获取跟进记录
    cursor.execute("""
        SELECT * FROM follow_records 
        WHERE customer_id = ? 
        ORDER BY created_at DESC
    """, (customer_id,))
    follow_records = [dict(r) for r in cursor.fetchall()]
    
    customer['follow_records'] = follow_records
    
    conn.close()
    
    return {"success": True, "data": customer}

@app.put("/api/customers/{customer_id}/update")
async def update_customer(
    customer_id: str,
    request: dict,
    token: str = Depends(check_token)
):
    """
    更新客户信息并同步到企业微信
    """
    try:
        remark = request.get('remark', '')
        owner_userid = request.get('owner_userid', '')
        tags = request.get('tags', [])  # 新增：标签ID列表
        
        # 1. 获取当前客户的标签（用于对比）
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT enterprise_tags FROM customers WHERE id = ?", (customer_id,))
        row = cursor.fetchone()
        
        old_tags = []
        if row and row['enterprise_tags']:
            try:
                old_tags_data = json.loads(row['enterprise_tags'])
                old_tags = [t['tag_id'] for t in old_tags_data]
            except:
                pass
        
        # 2. 更新本地数据库
        current_time = int(time.time())
        cursor.execute("""
            UPDATE customers 
            SET remark = ?, updated_at = ?
            WHERE id = ?
        """, (remark, current_time, customer_id))
        
        conn.commit()
        conn.close()
        
        # 3. 同步到企业微信
        if owner_userid:
            try:
                # 同步备注
                remark_result = wecom_client.update_customer_remark(
                    external_userid=customer_id,
                    userid=owner_userid,
                    remark=remark
                )
                
                if not remark_result:
                    return {
                        "success": False, 
                        "message": "同步备注到企业微信失败，但本地已保存"
                    }
                
                # 同步标签（如果有变化）
                if tags != old_tags:
                    # 计算要添加和移除的标签
                    add_tags = [t for t in tags if t not in old_tags]
                    remove_tags = [t for t in old_tags if t not in tags]
                    
                    if add_tags or remove_tags:
                        tags_result = wecom_client.update_customer_tags(
                            external_userid=customer_id,
                            userid=owner_userid,
                            add_tag=add_tags if add_tags else None,
                            remove_tag=remove_tags if remove_tags else None
                        )
                        
                        if not tags_result:
                            return {
                                "success": False,
                                "message": "同步标签到企业微信失败，但备注已同步"
                            }
                    
            except Exception as e:
                print(f"[更新客户] 同步企业微信失败: {e}")
                return {
                    "success": False,
                    "message": f"同步到企业微信失败: {str(e)}，但本地已保存"
                }
        
        return {"success": True, "message": "更新成功"}
        
    except Exception as e:
        print(f"[更新客户] 错误: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

@app.get("/api/employees")
async def get_employees(token: str = Depends(check_token)):
    """获取员工列表（含统计数据）"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM employees")
    rows = cursor.fetchall()
    employees = []
    
    for row in rows:
        emp = dict(row)
        userid = emp['id']
        
        # 1. 统计该员工创建的群数量
        cursor.execute("""
            SELECT COUNT(*) as group_count
            FROM customer_groups
            WHERE owner_userid = ?
        """, (userid,))
        group_result = cursor.fetchone()
        emp['group_count'] = group_result[0] if group_result else 0
        
        # 2. 统计该员工的好友（客户）数量
        cursor.execute("""
            SELECT COUNT(*) as customer_count
            FROM customers
            WHERE owner_userid = ?
        """, (userid,))
        customer_result = cursor.fetchone()
        emp['customer_count'] = customer_result[0] if customer_result else 0
        
        # 3. 统计该员工最近添加的客户数量（最近30天）
        thirty_days_ago = int(time.time()) - (30 * 24 * 3600)
        cursor.execute("""
            SELECT COUNT(*) as recent_customer_count
            FROM customers
            WHERE owner_userid = ? AND add_time >= ?
        """, (userid, thirty_days_ago))
        recent_result = cursor.fetchone()
        emp['recent_customer_count'] = recent_result[0] if recent_result else 0
        
        employees.append(emp)
    
    conn.close()
    
    return {"success": True, "data": employees}

@app.post("/api/sync/employees")
async def sync_employees(request: SyncRequest, token: str = Depends(check_token)):
    """同步员工数据"""
    try:
        # 更新配置
        if request.config:
            wecom_client.update_config(
                corp_id=request.config.get('corpid'),
                contact_secret=request.config.get('contact_secret'),
                customer_secret=request.config.get('customer_secret'),
                app_secret=request.config.get('app_secret'),
                agent_id=request.config.get('agentid')
            )
        
        users = wecom_client.get_user_list()
        
        print(f"[员工同步] 获取到 {len(users)} 个员工")
        
        conn = get_db()
        cursor = conn.cursor()
        current_time = int(time.time())
        
        import json
        synced_count = 0
        for user in users:
            userid = user['userid']
            
            # 打印前3个员工的详细信息（调试用）
            if synced_count < 3:
                print(f"[员工同步] 员工 {synced_count + 1}:")
                print(f"  - userid: {userid}")
                print(f"  - name: {user.get('name', '')}")
                print(f"  - mobile: {user.get('mobile', '无')}")
                print(f"  - email: {user.get('email', '无')}")
                print(f"  - biz_mail: {user.get('biz_mail', '无')}")
                print(f"  - position: {user.get('position', '无')}")
            
            cursor.execute("SELECT id FROM employees WHERE id = ?", (userid,))
            exists = cursor.fetchone()
            
            department_json = json.dumps(user.get('department', []))
            
            # 优先使用 mobile，如果没有则留空
            mobile = user.get('mobile', '')
            
            # 优先使用 email，如果没有则使用 biz_mail
            email = user.get('email', '') or user.get('biz_mail', '')
            
            if exists:
                cursor.execute("""
                    UPDATE employees SET
                        name = ?, avatar = ?, mobile = ?, email = ?,
                        department = ?, position = ?, status = ?,
                        updated_at = ?
                    WHERE id = ?
                """, (
                    user.get('name', ''),
                    user.get('avatar', ''),
                    mobile,
                    email,
                    department_json,
                    user.get('position', ''),
                    user.get('status', 1),
                    current_time,
                    userid
                ))
            else:
                cursor.execute("""
                    INSERT INTO employees (
                        id, name, avatar, mobile, email,
                        department, position, status,
                        created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    userid,
                    user.get('name', ''),
                    user.get('avatar', ''),
                    mobile,
                    email,
                    department_json,
                    user.get('position', ''),
                    user.get('status', 1),
                    current_time,
                    current_time
                ))
            
            synced_count += 1
        
        conn.commit()
        
        # 验证同步后的数据
        cursor.execute("""
            SELECT 
                COUNT(*) as total,
                SUM(CASE WHEN mobile IS NOT NULL AND mobile != '' THEN 1 ELSE 0 END) as has_mobile,
                SUM(CASE WHEN email IS NOT NULL AND email != '' THEN 1 ELSE 0 END) as has_email
            FROM employees
        """)
        stats = cursor.fetchone()
        print(f"[员工同步] 完成！")
        print(f"  - 总员工数: {stats[0]}")
        print(f"  - 有手机号: {stats[1]}")
        print(f"  - 有邮箱: {stats[2]}")
        
        conn.close()
        
        return {"success": True, "message": f"同步成功，共 {len(users)} 个员工（有手机号: {stats[1]}, 有邮箱: {stats[2]}）"}
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/tags")
async def get_tags(token: str = Depends(check_token)):
    """获取标签列表（按组分组）"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM customer_tags ORDER BY group_name, order_num")
    rows = cursor.fetchall()
    tags = [dict(row) for row in rows]
    conn.close()
    
    # 按组分组标签
    tag_groups_dict = {}
    for tag in tags:
        group_name = tag.get('group_name') or '未分组'
        if group_name not in tag_groups_dict:
            tag_groups_dict[group_name] = {
                'group_id': tag.get('id', '').split('_')[0] if '_' in tag.get('id', '') else group_name,
                'group_name': group_name,
                'tags': []
            }
        tag_groups_dict[group_name]['tags'].append({
            'tag_id': tag.get('id'),
            'tag_name': tag.get('name'),
            'color': tag.get('color', ''),
            'order_num': tag.get('order_num', 0)
        })
    
    tag_groups = list(tag_groups_dict.values())
    
    return {"success": True, "data": tag_groups}

@app.post("/api/sync/tags")
async def sync_tags(request: SyncRequest, token: str = Depends(check_token)):
    """同步标签数据"""
    try:
        # 更新配置
        if request.config:
            wecom_client.update_config(
                corp_id=request.config.get('corpid'),
                contact_secret=request.config.get('contact_secret'),
                customer_secret=request.config.get('customer_secret'),
                app_secret=request.config.get('app_secret'),
                agent_id=request.config.get('agentid')
            )
        
        tag_groups = wecom_client.get_corp_tag_list()
        
        conn = get_db()
        cursor = conn.cursor()
        current_time = int(time.time())
        
        tag_count = 0
        for group in tag_groups:
            group_name = group.get('group_name', '')
            tags = group.get('tag', [])
            
            for tag in tags:
                tag_id = tag.get('id')
                tag_name = tag.get('name')
                order_num = tag.get('order', 0)
                
                cursor.execute("SELECT id FROM customer_tags WHERE id = ?", (tag_id,))
                exists = cursor.fetchone()
                
                if exists:
                    cursor.execute("""
                        UPDATE customer_tags SET
                            name = ?, group_name = ?, order_num = ?, updated_at = ?
                        WHERE id = ?
                    """, (tag_name, group_name, order_num, current_time, tag_id))
                else:
                    cursor.execute("""
                        INSERT INTO customer_tags (id, name, group_name, order_num, created_at)
                        VALUES (?, ?, ?, ?, ?)
                    """, (tag_id, tag_name, group_name, order_num, current_time))
                
                tag_count += 1
        
        conn.commit()
        conn.close()
        
        return {"success": True, "message": f"同步成功，共 {tag_count} 个标签"}
        
    except Exception as e:
        return {"success": False, "message": str(e)}

# ========== 前端静态文件 ==========

# 挂载静态文件目录
app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/")
async def read_root():
    """返回前端首页（禁用缓存）"""
    frontend_file = Path("static/index.html")
    if frontend_file.exists():
        response = FileResponse(frontend_file)
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response
    return HTMLResponse(content="<h1>企业微信 CRM</h1><p>前端文件未找到</p>", status_code=404)

@app.get("/style.css")
async def get_css():
    """返回 CSS（禁用缓存）"""
    css_file = Path("static/style.css")
    if css_file.exists():
        response = FileResponse(css_file, media_type="text/css")
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response
    return HTMLResponse(content="/* CSS not found */", status_code=404)

@app.get("/script.js")
async def get_js():
    """返回 JavaScript（禁用缓存）"""
    js_file = Path("static/script.js")
    if js_file.exists():
        response = FileResponse(js_file, media_type="application/javascript")
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response
    return HTMLResponse(content="// JS not found", status_code=404)

@app.get("/group-tags.js")
async def get_group_tags_js():
    """返回客户群标签 JavaScript（禁用缓存）"""
    js_file = Path("static/group-tags.js")
    if js_file.exists():
        response = FileResponse(js_file, media_type="application/javascript")
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response
    return HTMLResponse(content="// group-tags.js not found", status_code=404)

@app.get("/group-tags-v2.js")
async def get_group_tags_v2_js():
    """返回客户群标签 JavaScript V2（禁用缓存）"""
    js_file = Path("static/group-tags-v2.js")
    if js_file.exists():
        response = FileResponse(js_file, media_type="application/javascript")
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response
    return HTMLResponse(content="// group-tags-v2.js not found", status_code=404)

# 测试页面路由
@app.get("/complete-test.html")
async def get_complete_test():
    """返回完整测试页面"""
    test_file = Path("static/complete-test.html")
    if test_file.exists():
        return FileResponse(test_file)
    return HTMLResponse(content="<h1>测试页面未找到</h1>", status_code=404)

@app.get("/test-upload-button.html")
async def get_test_upload():
    """返回上传按钮测试页面"""
    test_file = Path("static/test-upload-button.html")
    if test_file.exists():
        return FileResponse(test_file)
    return HTMLResponse(content="<h1>测试页面未找到</h1>", status_code=404)

@app.get("/debug.html")
async def get_debug():
    """返回调试页面"""
    debug_file = Path("static/debug.html")
    if debug_file.exists():
        return FileResponse(debug_file)
    return HTMLResponse(content="<h1>调试页面未找到</h1>", status_code=404)

# ========== 跟进记录 API ==========

class FollowRecordCreate(BaseModel):
    """添加跟进记录请求"""
    customer_id: str
    content: str
    follow_type: Optional[str] = "电话"

@app.post("/api/follow-records")
async def create_follow_record(record: FollowRecordCreate, token: str = Depends(check_token)):
    """添加跟进记录"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        import uuid
        record_id = str(uuid.uuid4())
        current_time = int(time.time() * 1000)
        
        # TODO: 获取当前用户ID（暂时使用默认值）
        follower_userid = "default_user"
        
        cursor.execute("""
            INSERT INTO follow_records (
                id, customer_id, follower_userid, content, follow_type, created_at
            ) VALUES (?, ?, ?, ?, ?, ?)
        """, (record_id, record.customer_id, follower_userid, record.content, record.follow_type, current_time))
        
        conn.commit()
        conn.close()
        
        return {"success": True, "message": "跟进记录添加成功", "id": record_id}
    except Exception as e:
        return {"success": False, "message": str(e)}

# ========== 导出到企业微信表格 API ==========

class ExportToSpreadsheetRequest(BaseModel):
    """导出到企业微信表格请求"""
    customer_ids: List[str]  # 客户ID列表
    doc_name: Optional[str] = None  # 表格名称
    admin_users: Optional[List[str]] = None  # 管理员列表

@app.post("/api/export/spreadsheet")
async def export_to_spreadsheet(
    request: ExportToSpreadsheetRequest,
    token: str = Depends(check_token),
    config: dict = Depends(get_wecom_config)
):
    """
    导出客户列表到企业微信表格
    """
    try:
        print(f"[API] 导出 {len(request.customer_ids)} 个客户到企业微信表格")
        
        # 更新企业微信配置
        wecom_client.update_config(
            corp_id=config.get('corpid'),
            contact_secret=config.get('contact_secret'),
            customer_secret=config.get('customer_secret'),
            app_secret=config.get('app_secret'),
            agent_id=config.get('agentid')
        )
        
        # 从数据库获取客户数据
        conn = get_db()
        cursor = conn.cursor()
        
        customers = []
        for customer_id in request.customer_ids:
            cursor.execute("SELECT * FROM customers WHERE id = ?", (customer_id,))
            row = cursor.fetchone()
            if row:
                customer = dict(row)
                
                # 解析 tags（JSON 字符串转列表）
                tags_str = customer.get('tags', '[]')
                if tags_str:
                    try:
                        customer['tags'] = json.loads(tags_str)
                    except:
                        customer['tags'] = []
                
                customers.append(customer)
        
        conn.close()
        
        if not customers:
            return {"success": False, "message": "未找到要导出的客户"}
        
        # 生成表格名称
        if not request.doc_name:
            from datetime import datetime
            doc_name = f"客户列表导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        else:
            doc_name = request.doc_name
        
        # 导出到企业微信表格
        result = wecom_client.export_customers_to_spreadsheet(
            customers=customers,
            doc_name=doc_name,
            admin_users=request.admin_users
        )
        
        return result
        
    except Exception as e:
        print(f"[API] 导出失败: {e}")
        return {"success": False, "message": str(e)}

# ========== 智能表格 API ==========

import json
import hashlib
import uuid
from fastapi import File, UploadFile

class CreateSpreadsheetRequest(BaseModel):
    """创建智能表格请求"""
    name: str
    fields: List[Dict]  # [{"name": "日期", "type": "date"}, ...]
    data: List[List]    # [["2024-01-23", "合约A", ...], ...]
    config: Optional[dict] = None  # 企业微信配置（可选）
    admin_users: Optional[List[str]] = None  # 管理员列表（可选）

class SyncSpreadsheetRequest(BaseModel):
    """同步表格请求"""
    config: Optional[dict] = None  # 企业微信配置（可选）

@app.post("/api/spreadsheet/upload")
async def upload_excel(
    file: UploadFile = File(...),
    token: str = Depends(check_token)
):
    """上传并解析 Excel 文件"""
    try:
        import openpyxl
        from io import BytesIO
        
        print(f"[API] 接收 Excel 文件: {file.filename}")
        
        # 保存文件
        upload_dir = Path("data/uploads")
        upload_dir.mkdir(parents=True, exist_ok=True)
        
        file_id = str(uuid.uuid4())
        file_path = upload_dir / f"{file_id}_{file.filename}"
        
        content = await file.read()
        with open(file_path, "wb") as f:
            f.write(content)
        
        # 解析 Excel
        workbook = openpyxl.load_workbook(BytesIO(content))
        sheet = workbook.active
        
        # 读取表头
        headers = []
        for cell in sheet[1]:
            if cell.value:
                headers.append(str(cell.value))
        
        # 读取数据
        data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if any(row):
                row_data = []
                for cell in row[:len(headers)]:  # 只读取表头对应的列
                    if cell is not None:
                        row_data.append(str(cell))
                    else:
                        row_data.append('')
                data.append(row_data)
        
        # 自动推断字段类型
        field_types = {}
        for i, header in enumerate(headers):
            field_types[header] = infer_field_type(header, [row[i] if i < len(row) else '' for row in data])
        
        print(f"[API] 解析完成: {len(headers)} 个字段, {len(data)} 行数据")
        
        return {
            "success": True,
            "file_id": file_id,
            "file_name": file.filename,
            "file_path": str(file_path),
            "sheet_name": sheet.title,
            "headers": headers,
            "data": data,
            "field_types": field_types,
            "row_count": len(data),
            "col_count": len(headers)
        }
        
    except Exception as e:
        print(f"[API] Excel 解析失败: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

def infer_field_type(field_name: str, sample_values: List[str]) -> str:
    """推断字段类型"""
    field_lower = field_name.lower()
    
    # 根据字段名推断
    if '日期' in field_name or '时间' in field_name or 'date' in field_lower or 'time' in field_lower:
        return 'datetime'
    elif '电话' in field_name or 'phone' in field_lower or 'tel' in field_lower:
        return 'phone'
    elif '邮箱' in field_name or 'email' in field_lower or 'mail' in field_lower:
        return 'email'
    elif '金额' in field_name or '价格' in field_name or 'amount' in field_lower or 'price' in field_lower:
        return 'money'
    else:
        # 根据数据推断
        non_empty = [v for v in sample_values if v and v.strip()]
        if not non_empty:
            return 'text'
        
        # 检查是否都是数字
        try:
            for v in non_empty[:10]:
                float(v)
            return 'number'
        except:
            return 'text'

@app.post("/api/spreadsheet/create")
async def create_spreadsheet(
    request: CreateSpreadsheetRequest,
    token: str = Depends(check_token)
):
    """从解析的数据创建企业微信表格"""
    try:
        print(f"[API] 创建智能表格: {request.name}")
        print(f"[配置] 接收到的配置: {request.config}")
        
        # 更新企业微信配置
        if request.config:
            corpid = request.config.get('corpid', '')
            app_secret = request.config.get('app_secret', '')
            print(f"[配置] Corp ID: {corpid}")
            print(f"[配置] App Secret 前10位: {app_secret[:10] if app_secret else '未提供'}...")
            
            wecom_client.update_config(
                corp_id=corpid,
                contact_secret=request.config.get('contact_secret'),
                customer_secret=request.config.get('customer_secret'),
                app_secret=app_secret,
                agent_id=request.config.get('agentid')
            )
        else:
            print(f"[配置] 警告：未提供配置信息")
        
        # 0. 尝试获取空间列表（如果需要 spaceid）
        print("[表格] 尝试获取空间列表...")
        space_result = wecom_client.get_space_list()
        spaceid = None
        if space_result.get('errcode') == 0 and space_result.get('space_list'):
            spaceid = space_result['space_list'][0].get('spaceid')
            print(f"[表格] 找到空间: {spaceid}")
        else:
            print(f"[表格] 未获取到空间列表，将不使用 spaceid")
        
        # 1. 创建表格
        create_result = wecom_client.create_spreadsheet(
            doc_name=request.name,
            admin_users=request.admin_users,
            spaceid=spaceid
        )
        
        if create_result.get('errcode') != 0:
            return {
                "success": False,
                "message": f"创建表格失败: {create_result.get('errmsg')}"
            }
        
        docid = create_result.get('docid')
        url = create_result.get('url')
        
        # 2. 准备数据（表头 + 数据）
        headers = [field['name'] for field in request.fields]
        values = [headers] + request.data
        
        # 3. 写入数据
        write_result = wecom_client.write_spreadsheet_data(docid, values)
        
        data_written = False
        warning_message = None
        
        if write_result.get('errcode') == 0:
            data_written = True
            print("[API] 数据写入成功")
        elif write_result.get('errcode') == -404:
            # 智能表格不支持 API 写入
            warning_message = "表格已创建，但智能表格暂不支持 API 写入数据，请手动导入数据"
            print(f"[API] {warning_message}")
        else:
            # 其他错误，记录但不阻止流程
            warning_message = f"表格已创建，但数据写入失败: {write_result.get('errmsg')}"
            print(f"[API] {warning_message}")
        
        # 4. 保存到数据库
        spreadsheet_id = str(uuid.uuid4())
        current_time = int(time.time())
        
        conn = get_db()
        cursor = conn.cursor()
        
        # 计算数据 hash
        data_hash = hashlib.md5(json.dumps(request.data, ensure_ascii=False).encode()).hexdigest()
        
        cursor.execute("""
            INSERT INTO smart_spreadsheets (
                id, docid, name, fields_config, row_count, col_count,
                created_at, updated_at, last_sync_at, version, data_hash, url
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            spreadsheet_id, docid, request.name,
            json.dumps(request.fields, ensure_ascii=False),
            len(request.data), len(headers),
            current_time, current_time, current_time, 1, data_hash, url
        ))
        
        # 保存数据
        for row_idx, row in enumerate(request.data):
            for col_idx, value in enumerate(row):
                data_id = str(uuid.uuid4())
                cursor.execute("""
                    INSERT INTO spreadsheet_data (
                        id, spreadsheet_id, row_index, col_index, col_name, value,
                        version, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    data_id, spreadsheet_id, row_idx + 2, col_idx,  # row_idx+2 因为第1行是表头
                    headers[col_idx], value, 1, current_time, current_time
                ))
        
        conn.commit()
        conn.close()
        
        print(f"[API] 表格创建成功: {spreadsheet_id}")
        
        # 构建返回消息
        if warning_message:
            message = warning_message
            success_type = "partial"  # 部分成功
        else:
            message = f"成功创建表格并写入 {len(request.data)} 行数据"
            success_type = "full"  # 完全成功
        
        return {
            "success": True,
            "success_type": success_type,
            "spreadsheet_id": spreadsheet_id,
            "docid": docid,
            "url": url,
            "message": message,
            "data_written": data_written
        }
        
    except Exception as e:
        print(f"[API] 创建表格失败: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

@app.get("/api/spreadsheet/list")
async def list_spreadsheets(token: str = Depends(check_token)):
    """获取智能表格列表"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT id, docid, name, row_count, col_count, created_at, updated_at, last_sync_at, url, status
            FROM smart_spreadsheets
            WHERE status = 'active'
            ORDER BY created_at DESC
        """)
        
        spreadsheets = []
        for row in cursor.fetchall():
            spreadsheets.append({
                "id": row[0],
                "docid": row[1],
                "name": row[2],
                "row_count": row[3],
                "col_count": row[4],
                "created_at": row[5],
                "updated_at": row[6],
                "last_sync_at": row[7],
                "url": row[8],
                "status": row[9]
            })
        
        conn.close()
        
        return {
            "success": True,
            "data": spreadsheets,
            "count": len(spreadsheets)
        }
        
    except Exception as e:
        print(f"[API] 获取表格列表失败: {e}")
        return {"success": False, "message": str(e)}

@app.get("/api/spreadsheet/{spreadsheet_id}")
async def get_spreadsheet(spreadsheet_id: str, token: str = Depends(check_token)):
    """获取表格详情"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        # 获取表格信息
        cursor.execute("""
            SELECT id, docid, name, file_name, fields_config, row_count, col_count,
                   created_at, updated_at, last_sync_at, version, data_hash, url, status
            FROM smart_spreadsheets
            WHERE id = ?
        """, (spreadsheet_id,))
        
        row = cursor.fetchone()
        if not row:
            return {"success": False, "message": "表格不存在"}
        
        spreadsheet = {
            "id": row[0],
            "docid": row[1],
            "name": row[2],
            "file_name": row[3],
            "fields": json.loads(row[4]) if row[4] else [],
            "row_count": row[5],
            "col_count": row[6],
            "created_at": row[7],
            "updated_at": row[8],
            "last_sync_at": row[9],
            "version": row[10],
            "data_hash": row[11],
            "url": row[12],
            "status": row[13]
        }
        
        # 获取数据
        cursor.execute("""
            SELECT row_index, col_index, col_name, value
            FROM spreadsheet_data
            WHERE spreadsheet_id = ?
            ORDER BY row_index, col_index
        """, (spreadsheet_id,))
        
        data_dict = {}
        for row in cursor.fetchall():
            row_idx = row[0]
            if row_idx not in data_dict:
                data_dict[row_idx] = []
            data_dict[row_idx].append(row[3])
        
        spreadsheet['data'] = [data_dict[i] for i in sorted(data_dict.keys())]
        
        conn.close()
        
        return {
            "success": True,
            "data": spreadsheet
        }
        
    except Exception as e:
        print(f"[API] 获取表格详情失败: {e}")
        return {"success": False, "message": str(e)}

@app.post("/api/spreadsheet/{spreadsheet_id}/sync")
async def sync_spreadsheet(
    spreadsheet_id: str,
    request: SyncSpreadsheetRequest,
    token: str = Depends(check_token)
):
    """同步表格数据"""
    try:
        print(f"[API] 同步表格: {spreadsheet_id}")
        
        # 更新企业微信配置
        if request.config:
            wecom_client.update_config(
                corp_id=request.config.get('corpid'),
                contact_secret=request.config.get('contact_secret'),
                customer_secret=request.config.get('customer_secret'),
                app_secret=request.config.get('app_secret'),
                agent_id=request.config.get('agentid')
            )
        
        # 1. 获取本地表格信息
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT docid, data_hash, version
            FROM smart_spreadsheets
            WHERE id = ?
        """, (spreadsheet_id,))
        
        row = cursor.fetchone()
        if not row:
            return {"success": False, "message": "表格不存在"}
        
        docid, old_hash, version = row[0], row[1], row[2]
        
        # 2. 读取企业微信表格数据
        read_result = wecom_client.read_spreadsheet_data(docid)
        
        if read_result.get('errcode') != 0:
            return {
                "success": False,
                "message": f"读取表格失败: {read_result.get('errmsg')}"
            }
        
        wecom_data = read_result.get('values', [])
        
        if not wecom_data or len(wecom_data) < 2:
            return {"success": False, "message": "表格数据为空"}
        
        # 3. 计算新的 hash
        data_without_header = wecom_data[1:]  # 去除表头
        new_hash = hashlib.md5(json.dumps(data_without_header, ensure_ascii=False).encode()).hexdigest()
        
        # 4. 对比 hash
        if new_hash == old_hash:
            print(f"[同步] 数据无变化")
            return {
                "success": True,
                "changed": False,
                "message": "数据已是最新"
            }
        
        # 5. 数据有变化，更新本地数据
        print(f"[同步] 检测到数据变化，开始更新")
        
        # 删除旧数据
        cursor.execute("DELETE FROM spreadsheet_data WHERE spreadsheet_id = ?", (spreadsheet_id,))
        
        # 插入新数据
        headers = wecom_data[0]
        current_time = int(time.time())
        new_version = version + 1
        
        for row_idx, row in enumerate(data_without_header):
            for col_idx, value in enumerate(row):
                data_id = str(uuid.uuid4())
                cursor.execute("""
                    INSERT INTO spreadsheet_data (
                        id, spreadsheet_id, row_index, col_index, col_name, value,
                        version, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    data_id, spreadsheet_id, row_idx + 2, col_idx,
                    headers[col_idx] if col_idx < len(headers) else f'Col{col_idx}',
                    value, new_version, current_time, current_time
                ))
        
        # 更新表格信息
        cursor.execute("""
            UPDATE smart_spreadsheets
            SET data_hash = ?, version = ?, last_sync_at = ?, updated_at = ?, row_count = ?
            WHERE id = ?
        """, (new_hash, new_version, current_time, current_time, len(data_without_header), spreadsheet_id))
        
        # 记录同步日志
        log_id = str(uuid.uuid4())
        cursor.execute("""
            INSERT INTO sync_logs (id, spreadsheet_id, sync_type, changes_count, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (log_id, spreadsheet_id, 'manual', len(data_without_header), 'success', current_time))
        
        conn.commit()
        conn.close()
        
        print(f"[同步] 同步完成，版本: {version} -> {new_version}")
        
        return {
            "success": True,
            "changed": True,
            "old_version": version,
            "new_version": new_version,
            "row_count": len(data_without_header),
            "message": f"同步成功，更新了 {len(data_without_header)} 行数据"
        }
        
    except Exception as e:
        print(f"[API] 同步失败: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

@app.delete("/api/spreadsheet/{spreadsheet_id}")
async def delete_spreadsheet(spreadsheet_id: str, token: str = Depends(check_token)):
    """删除表格（同时删除企业微信中的文档）"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        # 1. 获取表格的 docid
        cursor.execute("""
            SELECT docid FROM smart_spreadsheets
            WHERE id = ? AND status != 'deleted'
        """, (spreadsheet_id,))
        
        row = cursor.fetchone()
        if not row:
            conn.close()
            return {"success": False, "message": "表格不存在"}
        
        docid = row[0]
        
        # 2. 调用企业微信 API 删除文档
        print(f"[API] 正在删除企业微信文档: {docid}")
        delete_result = wecom_client.delete_spreadsheet(docid)
        
        # 3. 无论企业微信删除是否成功，都标记本地为已删除
        cursor.execute("""
            UPDATE smart_spreadsheets
            SET status = 'deleted', updated_at = ?
            WHERE id = ?
        """, (int(time.time()), spreadsheet_id))
        
        conn.commit()
        conn.close()
        
        if delete_result.get('errcode') == 0:
            print(f"[API] 成功删除企业微信文档: {docid}")
            return {"success": True, "message": "表格已删除（包括企业微信中的文档）"}
        else:
            # 企业微信删除失败，但本地已标记删除
            error_msg = delete_result.get('errmsg', '未知错误')
            print(f"[API] 企业微信文档删除失败: {error_msg}")
            return {
                "success": True,  # 本地删除成功
                "message": f"本地表格已删除，但企业微信文档删除失败: {error_msg}",
                "wecom_error": error_msg
            }
        
    except Exception as e:
        print(f"[API] 删除表格失败: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

# ========== 新增 API：字段模板管理 ==========

@app.get("/api/templates/list")
async def get_template_list(token: str = Depends(check_token)):
    """获取字段模板列表"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT id, name, data_type, description, is_system, created_at
            FROM field_templates
            ORDER BY is_system DESC, created_at DESC
        """)
        
        templates = []
        for row in cursor.fetchall():
            templates.append({
                "id": row[0],
                "name": row[1],
                "data_type": row[2],
                "description": row[3],
                "is_system": row[4] == 1,
                "created_at": row[5]
            })
        
        conn.close()
        return {"success": True, "data": templates}
        
    except Exception as e:
        print(f"[API] 获取模板列表失败: {e}")
        return {"success": False, "message": str(e)}

@app.get("/api/templates/{template_id}")
async def get_template_detail(template_id: str, token: str = Depends(check_token)):
    """获取模板详情（包含字段配置）"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT id, name, data_type, fields_config, description, is_system, created_at
            FROM field_templates
            WHERE id = ?
        """, (template_id,))
        
        row = cursor.fetchone()
        if not row:
            return {"success": False, "message": "模板不存在"}
        
        import json
        template = {
            "id": row[0],
            "name": row[1],
            "data_type": row[2],
            "fields": json.loads(row[3]),
            "description": row[4],
            "is_system": row[5] == 1,
            "created_at": row[6]
        }
        
        conn.close()
        return {"success": True, "data": template}
        
    except Exception as e:
        print(f"[API] 获取模板详情失败: {e}")
        return {"success": False, "message": str(e)}

# ========== 新增 API：手工创建表格 ==========

class FieldConfig(BaseModel):
    """字段配置"""
    wecom_name: str
    system_name: str
    type: str = "text"
    
    class Config:
        extra = 'allow'  # 允许额外字段

class ManualCreateTableRequest(BaseModel):
    """手工创建表格请求"""
    name: str
    data_type: str  # order/supplier/product
    data_scope: str = "global"  # global/supplier
    supplier_code: Optional[str] = None
    fields: List[FieldConfig]  # 字段列表
    field_mapping: Optional[Dict[str, str]] = None  # 字段映射
    sync_config: Optional[Dict[str, Any]] = None  # 同步配置
    config: Optional[Dict[str, Any]] = None
    
    class Config:
        extra = 'allow'  # 允许额外字段

# 调试接口：接收任意数据
from fastapi import Request as FastAPIRequest

@app.post("/api/spreadsheet/create-manual-debug")
async def create_spreadsheet_manual_debug(
    request: FastAPIRequest,
    token: str = Depends(check_token)
):
    """调试接口：查看接收到的原始数据"""
    import json
    body = await request.body()
    data = json.loads(body)
    
    print(f"[调试] ===== 接收到的原始数据 =====")
    print(json.dumps(data, indent=2, ensure_ascii=False))
    print(f"[调试] ================================")
    
    return {"success": True, "received_data": data}

# 使用原始请求数据，绕过 Pydantic 验证
@app.post("/api/spreadsheet/create-manual")
async def create_spreadsheet_manual(
    raw_request: FastAPIRequest,
    token: str = Depends(check_token)
):
    """手工创建智能表格（不上传Excel）- 使用原始数据"""
    import json
    
    try:
        # 直接解析原始数据
        body = await raw_request.body()
        request_data = json.loads(body)
        
        name = request_data.get('name')
        data_type = request_data.get('data_type', 'order')
        data_scope = request_data.get('data_scope', 'global')
        supplier_code = request_data.get('supplier_code')
        fields = request_data.get('fields', [])
        field_mapping = request_data.get('field_mapping')
        sync_config = request_data.get('sync_config')
        config = request_data.get('config')
        
        print(f"[表格] 手工创建: {name}, 数据类型: {data_type}, 字段数: {len(fields)}")
        
        # 验证字段数量
        if len(fields) > 150:
            return {"success": False, "message": "字段数量不能超过 150 个"}
        
        if len(fields) == 0:
            return {"success": False, "message": "至少需要添加 1 个字段"}
        
        # 更新企业微信配置
        if config:
            wecom_client.update_config(
                corp_id=config.get('corpid'),
                contact_secret=config.get('contact_secret'),
                customer_secret=config.get('customer_secret'),
                app_secret=config.get('app_secret'),
                agent_id=config.get('agentid')
            )
        
        # 1. 创建企业微信智能表格
        import hashlib
        
        # 获取 spaceid
        spaces_result = wecom_client.get_space_list()
        spaceid = None
        if spaces_result.get('errcode') == 0:
            spaces = spaces_result.get('space_list', [])
            if spaces:
                spaceid = spaces[0].get('spaceid')
                print(f"[表格] 使用空间: {spaceid}")
        
        # 创建表格
        result = wecom_client.create_spreadsheet(
            doc_name=name,
            admin_users=[],
            spaceid=spaceid
        )
        
        if result.get('errcode') != 0:
            return {
                "success": False,
                "message": f"创建表格失败: {result.get('errmsg')}"
            }
        
        docid = result.get('docid')
        url = result.get('url')
        sheet_id = None
        
        # 尝试获取 sheet_id（可选，如果失败也继续）
        try:
            sheet_result = wecom_client.get_spreadsheet_sheets(docid)
            if sheet_result.get('errcode') == 0:
                sheets = sheet_result.get('sheet_list', [])
                if sheets:
                    sheet_id = sheets[0].get('sheet_id')
                    print(f"[表格] 获取到 sheet_id: {sheet_id}")
        except Exception as e:
            print(f"[表格] ⚠️  获取 sheet_id 失败（可能该接口不支持）: {e}")
        
        print(f"[表格] 创建成功: docid={docid}, url={url}")
        
        # 2. 添加字段到企业微信表格（尝试不使用 sheet_id）
        headers = [field['wecom_name'] for field in fields]
        
        if len(headers) > 0:
            # 先尝试不带 sheet_id
            print(f"[表格] 尝试添加字段（不使用 sheet_id）...")
            add_result = wecom_client.add_spreadsheet_fields(
                docid=docid,
                sheet_id=None,  # 不使用 sheet_id
                headers=headers
            )
            
            if add_result.get('errcode') != 0:
                # 如果失败且有 sheet_id，再尝试使用 sheet_id
                if sheet_id:
                    print(f"[表格] 不使用 sheet_id 失败，尝试使用 sheet_id...")
                    add_result = wecom_client.add_spreadsheet_fields(
                        docid=docid,
                        sheet_id=sheet_id,
                        headers=headers
                    )
                
                if add_result.get('errcode') != 0:
                    return {
                        "success": False,
                        "message": f"添加字段失败: {add_result.get('errmsg')}"
                    }
            
            print(f"[表格] 已添加 {len(headers)} 个字段")
        else:
            print(f"[表格] ⚠️  没有字段需要添加")
        
        # 3. 保存表格信息到数据库
        import time
        import uuid
        
        now = int(time.time())
        spreadsheet_id = str(uuid.uuid4())
        
        # 构建字段映射
        if not field_mapping:
            field_mapping = {}
            # 自动生成映射
            for field in fields:
                field_mapping[field['wecom_name']] = field['system_name']
        
        # 构建同步配置
        if not sync_config:
            sync_config = {
                "auto_sync": False,
                "sync_direction": "none",
                "sync_interval": 0,
                "conflict_strategy": "system_first"
            }
        
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("""
            INSERT INTO smart_spreadsheets (
                id, docid, sheet_id, name, data_type, data_scope, supplier_code,
                fields_config, field_mapping, sync_config,
                row_count, col_count, created_at, updated_at, version, url, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            spreadsheet_id,
            docid,
            sheet_id,
            name,
            data_type,
            data_scope,
            supplier_code,
            json.dumps(fields, ensure_ascii=False),
            json.dumps(field_mapping, ensure_ascii=False),
            json.dumps(sync_config, ensure_ascii=False),
            0,  # row_count
            len(headers),  # col_count
            now,
            now,
            1,  # version
            url,
            'active'
        ))
        
        conn.commit()
        conn.close()
        
        # 4. 生成字段使用说明（根据 editable 属性）
        readonly_fields = [f for f in fields if not f.get('editable', False)]
        editable_fields = [f for f in fields if f.get('editable', False)]
        
        field_usage_note = {
            "total_fields": len(fields),
            "readonly_count": len(readonly_fields),
            "editable_count": len(editable_fields),
            "readonly_fields": [f['wecom_name'] for f in readonly_fields],
            "editable_fields": [f['wecom_name'] for f in editable_fields],
            "note": "⚠️ 由于企业微信 API 限制，无法通过接口设置字段为只读。请手动在企业微信中避免编辑标记为'只读'的字段。"
        }
        
        print(f"[表格] 📋 字段权限说明:")
        print(f"  - 总字段数: {len(fields)}")
        print(f"  - 只读字段: {len(readonly_fields)} 个")
        print(f"  - 可编辑字段: {len(editable_fields)} 个")
        if readonly_fields:
            print(f"  - 只读字段列表（前10个）: {', '.join([f['wecom_name'] for f in readonly_fields[:10]])}")
        
        print(f"[表格] 已保存到数据库: {spreadsheet_id}")
        
        return {
            "success": True,
            "message": "表格创建成功",
            "data": {
                "id": spreadsheet_id,
                "docid": docid,
                "sheet_id": sheet_id,
                "name": name,
                "url": url,
                "field_count": len(headers),
                "data_type": data_type,
                "data_scope": data_scope,
                "supplier_code": supplier_code,
                "field_usage_note": field_usage_note  # 新增：字段使用说明
            }
        }
        
    except Exception as e:
        print(f"[API] 手工创建表格失败: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

# ========== 新增 API：供应商列表 ==========

@app.get("/api/suppliers/list")
async def get_supplier_list(token: str = Depends(check_token)):
    """获取供应商列表（从自有系统）"""
    try:
        # TODO: 这里需要调用自有系统的接口获取供应商列表
        # 目前返回模拟数据
        suppliers = [
            {"code": "MALXIN", "name": "马尔新"},
            {"code": "SUP001", "name": "供应商A"},
            {"code": "SUP002", "name": "供应商B"},
            {"code": "SUP003", "name": "供应商C"}
        ]
        
        return {"success": True, "data": suppliers}
        
    except Exception as e:
        print(f"[API] 获取供应商列表失败: {e}")
        return {"success": False, "message": str(e)}

# ========== 新增 API：手工同步数据 ==========

@app.post("/api/spreadsheet/{spreadsheet_id}/sync-manual")
async def sync_spreadsheet_manual(
    spreadsheet_id: str,
    token: str = Depends(check_token)
):
    """手工同步表格数据（从自有系统拉取数据）"""
    try:
        print(f"[同步] 手工同步表格: {spreadsheet_id}")
        
        # 1. 获取表格配置
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT docid, sheet_id, name, data_type, data_scope, supplier_code, 
                   fields_config, field_mapping, sync_config
            FROM smart_spreadsheets
            WHERE id = ? AND status = 'active'
        """, (spreadsheet_id,))
        
        row = cursor.fetchone()
        if not row:
            return {"success": False, "message": "表格不存在或已被删除"}
        
        import json
        docid = row[0]
        sheet_id = row[1]
        name = row[2]
        data_type = row[3]
        data_scope = row[4]
        supplier_code = row[5]
        fields_config = json.loads(row[6])
        field_mapping = json.loads(row[7])
        sync_config = json.loads(row[8])
        
        print(f"[同步] 表格配置: 数据类型={data_type}, 数据范围={data_scope}, 供应商={supplier_code}")
        
        # 2. 从自有系统获取数据
        # TODO: 调用自有系统接口获取数据
        # 这里返回模拟数据
        
        print(f"[同步] ⚠️  自有系统接口尚未对接，请提供以下接口：")
        print(f"   1. 订单列表: GET /api/orders/list?scope={data_scope}&supplier={supplier_code}")
        print(f"   2. 代理商列表: GET /api/suppliers/list?scope={data_scope}")
        print(f"   3. 产品列表: GET /api/products/list?scope={data_scope}&supplier={supplier_code}")
        
        return {
            "success": False,
            "message": "自有系统接口尚未对接，请联系技术团队提供数据接口",
            "required_apis": [
                {"name": "订单列表", "url": f"GET /api/orders/list?scope={data_scope}&supplier={supplier_code}"},
                {"name": "代理商列表", "url": "GET /api/suppliers/list?scope={data_scope}"},
                {"name": "产品列表", "url": f"GET /api/products/list?scope={data_scope}&supplier={supplier_code}"}
            ]
        }
        
    except Exception as e:
        print(f"[API] 手工同步失败: {e}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}

# ========== 客户导出 API ==========

class ExportCustomersRequest(BaseModel):
    """导出客户请求"""
    customer_ids: Optional[List[str]] = None  # 要导出的客户ID列表
    filters: Optional[dict] = None  # 筛选条件
    include_avatar: bool = True  # 是否包含头像

@app.post("/api/customers/export")
async def export_customers(request: ExportCustomersRequest, token: str = Depends(check_token)):
    """
    导出客户数据为Excel文件
    """
    try:
        print(f"[导出] 开始导出客户数据...")
        
        exporter = CustomerExporter(DB_PATH)
        
        # 导出Excel
        excel_bytes = exporter.export_to_excel(
            customer_ids=request.customer_ids,
            filters=request.filters,
            include_avatar=request.include_avatar
        )
        
        # 生成文件名
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"客户数据_{timestamp}.xlsx"
        
        print(f"[导出] 导出完成，文件大小: {len(excel_bytes)} 字节")
        
        from fastapi.responses import Response
        return Response(
            content=excel_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )
    
    except Exception as e:
        print(f"[导出] 导出失败: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

# ========== 定时同步任务 ==========

# import threading
# from datetime import time as dt_time
# import schedule

# 定时任务临时禁用
def sync_task_disabled():
    """定时同步任务"""
    try:
        print(f"\n{'='*60}")
        print(f"[定时任务] 开始自动同步 - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"{'='*60}")
        
        # 记录开始时间
        start_time = time.time()
        
        # 同步通讯录
        print("[定时任务] 1/2 同步通讯录...")
        try:
            users = wecom_client.get_user_list()
            # 保存到数据库（复用同步逻辑）
            conn = get_db()
            cursor = conn.cursor()
            current_time = int(time.time())
            
            import json
            for user in users:
                userid = user['userid']
                cursor.execute("SELECT id FROM employees WHERE id = ?", (userid,))
                exists = cursor.fetchone()
                
                department_json = json.dumps(user.get('department', []))
                
                if exists:
                    cursor.execute("""
                        UPDATE employees SET
                            name = ?, avatar = ?, mobile = ?, email = ?,
                            department = ?, position = ?, status = ?,
                            updated_at = ?
                        WHERE id = ?
                    """, (
                        user.get('name', ''),
                        user.get('avatar', ''),
                        user.get('mobile', ''),
                        user.get('email', ''),
                        department_json,
                        user.get('position', ''),
                        user.get('status', 1),
                        current_time,
                        userid
                    ))
                else:
                    cursor.execute("""
                        INSERT INTO employees (
                            id, name, avatar, mobile, email,
                            department, position, status,
                            created_at, updated_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        userid,
                        user.get('name', ''),
                        user.get('avatar', ''),
                        user.get('mobile', ''),
                        user.get('email', ''),
                        department_json,
                        user.get('position', ''),
                        user.get('status', 1),
                        current_time,
                        current_time
                    ))
            
            conn.commit()
            conn.close()
            print(f"[定时任务] ✅ 通讯录同步成功，共 {len(users)} 个员工")
        except Exception as e:
            print(f"[定时任务] ❌ 通讯录同步失败: {e}")
        
        # 同步客户
        print("[定时任务] 2/2 同步客户...")
        try:
            customers = wecom_client.sync_all_customers()
            
            if customers:
                conn = get_db()
                cursor = conn.cursor()
                current_time = int(time.time())
                
                saved_count = 0
                for customer in customers:
                    external_userid = customer.get('external_userid')
                    if not external_userid:
                        continue
                    
                    cursor.execute("SELECT id FROM customers WHERE id = ?", (external_userid,))
                    exists = cursor.fetchone()
                    
                    # 准备标签数据
                    import json
                    tags_json = json.dumps([tag.get('tag_name', '') for tag in customer.get('tags', [])], ensure_ascii=False)
                    
                    # 分类标签
                    enterprise_tags = []
                    personal_tags = []
                    rule_tags = []
                    
                    for tag in customer.get('tags', []):
                        tag_type = tag.get('type', 1)
                        tag_data = {
                            'tag_id': tag.get('tag_id', ''),
                            'tag_name': tag.get('tag_name', ''),
                            'group_name': tag.get('group_name', '')
                        }
                        if tag_type == 1:
                            enterprise_tags.append(tag_data)
                        elif tag_type == 2:
                            personal_tags.append(tag_data)
                        elif tag_type == 3:
                            rule_tags.append(tag_data)
                    
                    enterprise_tags_json = json.dumps(enterprise_tags, ensure_ascii=False)
                    personal_tags_json = json.dumps(personal_tags, ensure_ascii=False)
                    rule_tags_json = json.dumps(rule_tags, ensure_ascii=False)
                    
                    if exists:
                        cursor.execute("""
                            UPDATE customers SET
                                name = ?, avatar = ?, gender = ?, type = ?,
                                corp_name = ?, position = ?,
                                owner_userid = ?, owner_name = ?,
                                add_time = ?, tags = ?, remark = ?,
                                description = ?, add_way = ?, im_status = ?, state = ?,
                                remark_mobiles = ?, remark_corp_name = ?,
                                enterprise_tags = ?, personal_tags = ?, rule_tags = ?,
                                unionid = ?,
                                updated_at = ?
                            WHERE id = ?
                        """, (
                            customer.get('name', ''),
                            customer.get('avatar', ''),
                            customer.get('gender', 0),
                            customer.get('type', 1),
                            customer.get('corp_name', ''),
                            customer.get('position', ''),
                            customer.get('owner_userid', ''),
                            customer.get('owner_name', ''),
                            customer.get('add_time', 0),
                            tags_json,
                            customer.get('remark', ''),
                            customer.get('description', ''),
                            customer.get('add_way', 0),
                            customer.get('im_status', 0),
                            customer.get('state', ''),
                            json.dumps(customer.get('remark_mobiles', []), ensure_ascii=False),
                            customer.get('remark_corp_name', ''),
                            enterprise_tags_json,
                            personal_tags_json,
                            rule_tags_json,
                            customer.get('unionid', ''),
                            current_time,
                            external_userid
                        ))
                    else:
                        cursor.execute("""
                            INSERT INTO customers (
                                id, name, avatar, gender, type,
                                corp_name, position,
                                owner_userid, owner_name,
                                add_time, tags, remark,
                                description, add_way, im_status, state,
                                remark_mobiles, remark_corp_name,
                                enterprise_tags, personal_tags, rule_tags,
                                unionid,
                                created_at, updated_at
                            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (
                            external_userid,
                            customer.get('name', ''),
                            customer.get('avatar', ''),
                            customer.get('gender', 0),
                            customer.get('type', 1),
                            customer.get('corp_name', ''),
                            customer.get('position', ''),
                            customer.get('owner_userid', ''),
                            customer.get('owner_name', ''),
                            customer.get('add_time', 0),
                            tags_json,
                            customer.get('remark', ''),
                            customer.get('description', ''),
                            customer.get('add_way', 0),
                            customer.get('im_status', 0),
                            customer.get('state', ''),
                            json.dumps(customer.get('remark_mobiles', []), ensure_ascii=False),
                            customer.get('remark_corp_name', ''),
                            enterprise_tags_json,
                            personal_tags_json,
                            rule_tags_json,
                            customer.get('unionid', ''),
                            current_time,
                            current_time
                        ))
                    
                    saved_count += 1
                
                conn.commit()
                conn.close()
                print(f"[定时任务] ✅ 客户同步成功，共 {saved_count} 个客户")
            else:
                print("[定时任务] ⚠️ 未获取到客户数据")
        except Exception as e:
            print(f"[定时任务] ❌ 客户同步失败: {e}")
            import traceback
            traceback.print_exc()
        
        # 记录耗时
        duration = time.time() - start_time
        print(f"{'='*60}")
        print(f"[定时任务] 自动同步完成，耗时 {duration:.2f} 秒")
        print(f"{'='*60}\n")
        
        # 记录到同步日志表
        try:
            conn = get_db()
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO sync_logs (
                    sync_type, sync_time, success, duration_seconds, trigger_type
                ) VALUES (?, ?, ?, ?, ?)
            """, ('auto_full', int(time.time()), 1, duration, 'scheduled'))
            conn.commit()
            conn.close()
        except:
            pass
    
    except Exception as e:
        print(f"[定时任务] 执行失败: {e}")
        import traceback
        traceback.print_exc()

# def run_scheduler():
#     """运行定时任务调度器"""
#     while True:
#         schedule.run_pending()
#         time.sleep(60)  # 每分钟检查一次

def start_scheduled_sync():
    """启动定时同步"""
    # 临时禁用定时任务
    print(f"⚠️  定时同步已禁用（缺少schedule模块）")
    # # 每天凌晨00:00执行同步
    # schedule.every().day.at("00:00").do(sync_task)
    # 
    # # 启动调度器线程
    # scheduler_thread = threading.Thread(target=run_scheduler, daemon=True)
    # scheduler_thread.start()
    # 
    # print(f"⏰ 定时同步已启动：每天 00:00 自动同步")

# ========== 客户群管理 API ==========

@app.get("/api/customer-groups")
async def get_customer_groups(
    token: str = Depends(check_token),
    page: int = Query(1, gt=0),
    limit: int = Query(20, gt=0, le=100),
    search: Optional[str] = None,
    owner_userid: Optional[str] = None,
    group_type: Optional[str] = None,
    date_start: Optional[str] = None,
    date_end: Optional[str] = None,
    tag_id: Optional[str] = None
):
    """获取客户群列表"""
    conn = get_db()
    cursor = conn.cursor()
    
    # 构建查询条件
    where_clauses = []
    params = []
    
    if search:
        where_clauses.append("(name LIKE ? OR owner_name LIKE ?)")
        params.extend([f"%{search}%", f"%{search}%"])
    
    if owner_userid:
        where_clauses.append("owner_userid = ?")
        params.append(owner_userid)
    
    if group_type:
        where_clauses.append("group_type = ?")
        params.append(group_type)
    
    if date_start:
        where_clauses.append("create_time >= ?")
        params.append(int(datetime.fromisoformat(date_start).timestamp()))
    
    if date_end:
        where_clauses.append("create_time <= ?")
        params.append(int(datetime.fromisoformat(date_end + " 23:59:59").timestamp()))
    
    # 标签筛选
    if tag_id:
        where_clauses.append("""
            chat_id IN (
                SELECT chat_id FROM group_chat_tag_relations WHERE tag_id = ?
            )
        """)
        params.append(tag_id)
    
    where_sql = " AND ".join(where_clauses) if where_clauses else "1=1"
    
    # 查询总数
    count_sql = f"SELECT COUNT(*) FROM customer_groups WHERE {where_sql}"
    total = cursor.execute(count_sql, params).fetchone()[0]
    
    # 查询数据 - 显式指定字段名，避免字段顺序问题
    offset = (page - 1) * limit
    query_sql = f"""
        SELECT 
            chat_id, name, owner_userid, owner_name, notice, member_count,
            external_member_count, internal_member_count, admin_list, group_type,
            status, version, create_time, last_sync_time
        FROM customer_groups
        WHERE {where_sql}
        ORDER BY create_time DESC
        LIMIT ? OFFSET ?
    """
    
    rows = cursor.execute(query_sql, params + [limit, offset]).fetchall()
    
    groups = []
    for row in rows:
        chat_id = row[0]
        
        # 查询该群的标签（直接从关联表读取 tag_name）
        cursor.execute('''
            SELECT r.tag_id, r.tag_name
            FROM group_chat_tag_relations r
            WHERE r.chat_id = ?
            ORDER BY r.created_at DESC
        ''', (chat_id,))
        
        tags = []
        for tag_row in cursor.fetchall():
            tags.append({
                'tag_id': tag_row[0],
                'tag_name': tag_row[1]
            })
        
        groups.append({
            'chat_id': chat_id,
            'name': row[1],
            'owner_userid': row[2],
            'owner_name': row[3],
            'notice': row[4],
            'member_count': row[5],
            'external_member_count': row[6],
            'internal_member_count': row[7],
            'admin_list': row[8],
            'group_type': row[9],
            'status': row[10],
            'version': row[11],
            'create_time': row[12],
            'last_sync_time': row[13],
            'tags': tags  # 新增：标签列表
        })
    
    conn.close()
    
    return {
        "success": True,
        "data": groups,
        "total": total,
        "page": page,
        "limit": limit
    }

@app.post("/api/sync/customer-groups")
async def sync_customer_groups(request: SyncRequest, token: str = Depends(check_token)):
    """同步客户群数据（异步任务）"""
    print("[开始同步客户群]")
    
    # 更新企业微信配置
    if request.config:
        wecom_client.update_config(
            corp_id=request.config.get('corpid', ''),
            contact_secret=request.config.get('contact_secret', ''),
            customer_secret=request.config.get('customer_secret', ''),
            app_secret=request.config.get('app_secret', ''),
            agent_id=request.config.get('agentid', '')
        )
    
    # 使用同步服务创建异步任务
    task_id = sync_service.sync_customer_groups_async()
    
    return {
        "success": True,
        "task_id": task_id,
        "message": "客户群同步任务已启动"
    }

@app.get("/api/sync/customer-groups/status/{task_id}")
async def get_group_sync_status(task_id: str, token: str = Depends(check_token)):
    """获取客户群同步任务状态"""
    status = sync_service.get_task_status(task_id)
    if status:
        return {"success": True, "data": status}
    else:
        return {"success": False, "message": "任务不存在"}

@app.post("/api/sync/customer-groups/cancel/{task_id}")
async def cancel_group_sync(task_id: str, token: str = Depends(check_token)):
    """取消客户群同步任务"""
    success = sync_service.cancel_task(task_id)
    if success:
        return {"success": True, "message": "任务已取消"}
    else:
        return {"success": False, "message": "取消失败或任务不存在"}

# ========== 客户画像 API ==========

@app.get("/api/customer-portrait/tag-stats")
async def get_tag_stats(token: str = Depends(check_token)):
    """获取标签统计数据 - 修复版"""
    import json
    from collections import defaultdict
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        # 获取所有客户的企业标签
        cursor.execute("""
            SELECT id, name, enterprise_tags 
            FROM customers 
            WHERE enterprise_tags IS NOT NULL AND enterprise_tags != '' AND enterprise_tags != '[]'
        """)
        
        # 标签名称 -> 客户列表
        tag_customers = defaultdict(list)
        
        # 统计每个标签的客户数（一个客户有多个相同标签名，只算一次）
        for customer_id, customer_name, tags_str in cursor.fetchall():
            try:
                tags = json.loads(tags_str)
                if not isinstance(tags, list):
                    continue
                
                # 当前客户的所有标签名（去重）
                customer_tag_names = set()
                for tag in tags:
                    if isinstance(tag, dict):
                        tag_name = tag.get('tag_name', '').strip()
                        if tag_name:
                            customer_tag_names.add(tag_name)
                
                # 将客户添加到每个标签的列表中
                for tag_name in customer_tag_names:
                    tag_customers[tag_name].append({
                        'id': customer_id,
                        'name': customer_name or customer_id
                    })
                    
            except Exception as e:
                print(f"解析客户 {customer_id} 标签失败: {e}")
                continue
        
        # 计算每个标签的客户数量
        tag_stats = {}
        for tag_name, customers in tag_customers.items():
            tag_stats[tag_name] = len(customers)
        
        # 构建重点标签数据（固定的6个卡片）
        key_tags = {
            "用户标签": 0,
            "代理商": 0,
            "合伙人": 0,
            "供应商": 0,
            "同行": 0,
            "原有老代理": 0
        }
        
        # 匹配规则（严格匹配，优先匹配更具体的标签）
        for tag_name, count in tag_stats.items():
            # 用户标签
            if tag_name in ['用户', '客户', '用户标签']:
                key_tags["用户标签"] += count
            # 原有老代理商（优先匹配，避免和"代理商"冲突）
            elif '原有老代理商' in tag_name or '原有老代理' in tag_name or '老代理商' in tag_name:
                key_tags["原有老代理"] += count
            # 代理商（排除原有老代理商）
            elif tag_name in ['代理商', '代理'] and '原有' not in tag_name and '老' not in tag_name:
                key_tags["代理商"] += count
            # 合伙人
            elif '合伙人' in tag_name:
                key_tags["合伙人"] += count
            # 供应商
            elif '供应商' in tag_name:
                key_tags["供应商"] += count
            # 同行
            elif '同行' in tag_name or tag_name == '同行':
                key_tags["同行"] += count
        
        conn.close()
        
        return {
            "success": True,
            "data": {
                "key_tags": key_tags,
                "all_tag_stats": tag_stats,
                "tag_customers": {k: v for k, v in tag_customers.items()}
            }
        }
        
    except Exception as e:
        print(f"获取标签统计失败: {e}")
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/province-stats")
async def get_province_stats(token: str = Depends(check_token)):
    """获取省份统计数据（只显示有人的省份）"""
    import json
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT enterprise_tags 
            FROM customers 
            WHERE enterprise_tags IS NOT NULL AND enterprise_tags != '' AND enterprise_tags != '[]'
        """)
        
        province_counts = {}
        
        for row in cursor.fetchall():
            try:
                tags = json.loads(row[0])
                if isinstance(tags, list):
                    for tag in tags:
                        if isinstance(tag, dict):
                            group_name = tag.get('group_name', '')
                            tag_name = tag.get('tag_name', '')
                            
                            # 识别省份标签（假设省份在 group_name 或 tag_name 中）
                            # 可以根据实际情况调整匹配逻辑
                            if '省' in tag_name or '市' in tag_name or any(p in tag_name for p in [
                                '北京', '上海', '天津', '重庆', '河北', '山西', '辽宁', '吉林', '黑龙江',
                                '江苏', '浙江', '安徽', '福建', '江西', '山东', '河南', '湖北', '湖南',
                                '广东', '海南', '四川', '贵州', '云南', '陕西', '甘肃', '青海', '台湾',
                                '内蒙古', '广西', '西藏', '宁夏', '新疆', '香港', '澳门'
                            ]):
                                province_counts[tag_name] = province_counts.get(tag_name, 0) + 1
            except:
                continue
        
        # 只返回有人的省份，按人数排序
        sorted_provinces = sorted(province_counts.items(), key=lambda x: x[1], reverse=True)
        
        conn.close()
        
        return {
            "success": True,
            "data": dict(sorted_provinces[:30])  # 最多显示前30个省份
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/add-way-stats")
async def get_add_way_stats(token: str = Depends(check_token)):
    """获取添加方式统计数据"""
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT add_way, COUNT(*) as count
            FROM customers
            WHERE add_way IS NOT NULL
            GROUP BY add_way
            ORDER BY count DESC
        """)
        
        # 企业微信添加方式映射
        add_way_map = {
            0: "未知来源",
            1: "扫描二维码",
            2: "搜索手机号",
            3: "名片分享",
            4: "群聊",
            5: "手机通讯录",
            6: "微信联系人",
            7: "添加好友申请",
            8: "第三方应用",
            9: "搜索邮箱",
            16: "其他",
            22: "其他",
            24: "其他",
            201: "内部成员共享",
            202: "管理员分配"
        }
        
        results = []
        total = 0
        
        for row in cursor.fetchall():
            add_way = row[0]
            count = row[1]
            total += count
            
            results.append({
                "way": add_way_map.get(add_way, f"未知({add_way})"),
                "count": count
            })
        
        # 计算百分比
        for item in results:
            item["percentage"] = round(item["count"] / total * 100, 1) if total > 0 else 0
        
        conn.close()
        
        return {
            "success": True,
            "data": {
                "items": results,
                "total": total
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/gender-stats")
async def get_gender_stats(token: str = Depends(check_token)):
    """获取性别统计数据"""
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT gender, COUNT(*) as count
            FROM customers
            WHERE gender IS NOT NULL
            GROUP BY gender
            ORDER BY gender
        """)
        
        gender_map = {0: "未知", 1: "男", 2: "女"}
        
        results = []
        total = 0
        
        for row in cursor.fetchall():
            gender = row[0]
            count = row[1]
            total += count
            
            results.append({
                "gender": gender_map.get(gender, "未知"),
                "count": count
            })
        
        # 计算百分比
        for item in results:
            item["percentage"] = round(item["count"] / total * 100, 1) if total > 0 else 0
        
        conn.close()
        
        return {
            "success": True,
            "data": {
                "items": results,
                "total": total
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/time-stats")
async def get_time_stats(token: str = Depends(check_token)):
    """获取时间维度的客户新增统计"""
    from datetime import datetime, timedelta
    
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        now = datetime.now()
        
        # 今天的时间范围（0点到当前时间）
        today_start = datetime(now.year, now.month, now.day)
        
        # 昨天的时间范围（昨天0点到23:59:59）
        yesterday_start = today_start - timedelta(days=1)
        
        # 本周的时间范围（本周一0点到当前时间）
        weekday = now.weekday()  # 0=周一, 6=周日
        week_start = today_start - timedelta(days=weekday)
        
        # 上周的时间范围（上周一0点到上周日23:59:59）
        last_week_start = week_start - timedelta(days=7)
        last_week_end = week_start
        
        # 本月的时间范围（本月1号0点到当前时间）
        month_start = datetime(now.year, now.month, 1)
        
        # 上月的时间范围（上月1号到上月最后一天）
        if now.month == 1:
            last_month_start = datetime(now.year - 1, 12, 1)
            last_month_end = datetime(now.year, 1, 1)
        else:
            last_month_start = datetime(now.year, now.month - 1, 1)
            last_month_end = month_start
        
        # 查询各时间段的数据
        def count_customers(start_time, end_time=None):
            start_ts = int(start_time.timestamp())
            if end_time:
                end_ts = int(end_time.timestamp())
                cursor.execute("""
                    SELECT COUNT(*) FROM customers
                    WHERE add_time >= ? AND add_time < ?
                """, (start_ts, end_ts))
            else:
                # 如果没有结束时间，查询从开始时间到现在
                cursor.execute("""
                    SELECT COUNT(*) FROM customers
                    WHERE add_time >= ?
                """, (start_ts,))
            return cursor.fetchone()[0]
        
        stats = {
            "today": count_customers(today_start),
            "yesterday": count_customers(yesterday_start, today_start),
            "this_week": count_customers(week_start),
            "last_week": count_customers(last_week_start, last_week_end),
            "this_month": count_customers(month_start),
            "last_month": count_customers(last_month_start, last_month_end)
        }
        
        conn.close()
        
        return {
            "success": True,
            "data": stats
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/customers-by-tag")
async def get_customers_by_tag(
    tag_type: str = Query(..., description="标签类型：user/agent/partner/supplier/peer/old-agent"),
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    token: str = Depends(check_token)
):
    """根据标签类型获取客户列表"""
    import json
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        # 标签关键词映射
        tag_keywords = {
            "user": ["用户", "客户"],
            "agent": ["代理商", "代理"],
            "partner": ["合伙人"],
            "supplier": ["供应商"],
            "peer": ["同行", "上游同行"],
            "old-agent": ["原有老代理", "历史-代理"]
        }
        
        keywords = tag_keywords.get(tag_type, [])
        if not keywords:
            return {"success": False, "message": "无效的标签类型"}
        
        # 查询符合条件的客户
        cursor.execute("""
            SELECT id, name, enterprise_tags, avatar, gender, add_time, owner_userid, owner_name, remark
            FROM customers
            WHERE enterprise_tags IS NOT NULL AND enterprise_tags != '' AND enterprise_tags != '[]'
        """)
        
        matched_customers = []
        
        for row in cursor.fetchall():
            try:
                tags = json.loads(row[2]) if row[2] else []
                if isinstance(tags, list):
                    # 检查是否有匹配的标签
                    has_match = False
                    customer_tags = []
                    
                    for tag in tags:
                        if isinstance(tag, dict):
                            tag_name = tag.get('tag_name', '')
                            group_name = tag.get('group_name', '')
                            
                            # 收集所有标签
                            if tag_name:
                                customer_tags.append(tag_name)
                            
                            # 检查是否匹配
                            for keyword in keywords:
                                if keyword in tag_name or keyword in group_name:
                                    has_match = True
                                    break
                            
                            if has_match:
                                break
                    
                    if has_match:
                        matched_customers.append({
                            "id": row[0],
                            "name": row[1],
                            "tags": customer_tags,
                            "avatar": row[3] or "",
                            "gender": row[4] or 0,
                            "add_time": row[5] or 0,
                            "owner_userid": row[6] or "",
                            "owner_name": row[7] or "",
                            "remark": row[8] or ""
                        })
            except:
                continue
        
        # 分页
        total = len(matched_customers)
        start = (page - 1) * limit
        end = start + limit
        page_data = matched_customers[start:end]
        
        conn.close()
        
        return {
            "success": True,
            "data": {
                "customers": page_data,
                "total": total,
                "page": page,
                "limit": limit,
                "total_pages": (total + limit - 1) // limit
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/customers-by-filter")
async def get_customers_by_filter(
    filter_type: str = Query(..., description="筛选类型：province/add_way/gender"),
    filter_value: str = Query(..., description="筛选值"),
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    token: str = Depends(check_token)
):
    """根据省份/添加方式/性别筛选客户列表"""
    import json
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        matched_customers = []
        
        if filter_type == "province":
            # 按省份筛选
            cursor.execute("""
                SELECT id, name, enterprise_tags, avatar, gender, add_time, owner_userid, owner_name, remark
                FROM customers
                WHERE enterprise_tags IS NOT NULL AND enterprise_tags != '' AND enterprise_tags != '[]'
            """)
            
            for row in cursor.fetchall():
                try:
                    tags = json.loads(row[2]) if row[2] else []
                    if isinstance(tags, list):
                        # 检查是否有匹配的省份标签
                        for tag in tags:
                            if isinstance(tag, dict):
                                tag_name = tag.get('tag_name', '')
                                if filter_value in tag_name or tag_name == filter_value:
                                    customer_tags = [t.get('tag_name', '') for t in tags if isinstance(t, dict)][:5]
                                    matched_customers.append({
                                        "id": row[0],
                                        "name": row[1],
                                        "avatar": row[3] or "",
                                        "gender": row[4] or 0,
                                        "add_time": row[5] or 0,
                                        "owner_userid": row[6] or "",
                                        "owner_name": row[7] or "-",
                                        "remark": row[8] or "",
                                        "tags": customer_tags
                                    })
                                    break
                except:
                    continue
        
        elif filter_type == "add_way":
            # 按添加方式筛选
            add_way_map = {
                "未知来源": 0,
                "扫描二维码": 1,
                "搜索手机号": 2,
                "名片分享": 3,
                "群聊": 4,
                "手机通讯录": 5,
                "微信联系人": 6,
                "添加好友申请": 7,
                "第三方应用": 8,
                "搜索邮箱": 9,
                "内部成员共享": 201,
                "管理员分配": 202,
                "其他": [16, 22, 24]
            }
            
            # 找到对应的add_way值
            add_way_value = None
            for key, value in add_way_map.items():
                if filter_value == key:
                    add_way_value = value
                    break
            
            if add_way_value is not None:
                if isinstance(add_way_value, list):
                    # 处理"其他"情况
                    placeholders = ','.join(['?'] * len(add_way_value))
                    cursor.execute(f"""
                        SELECT id, name, enterprise_tags, avatar, gender, add_time, owner_userid, owner_name, remark
                        FROM customers
                        WHERE add_way IN ({placeholders})
                    """, add_way_value)
                else:
                    cursor.execute("""
                        SELECT id, name, enterprise_tags, avatar, gender, add_time, owner_userid, owner_name, remark
                        FROM customers
                        WHERE add_way = ?
                    """, (add_way_value,))
                
                for row in cursor.fetchall():
                    try:
                        tags = json.loads(row[2]) if row[2] else []
                        customer_tags = [t.get('tag_name', '') for t in tags if isinstance(t, dict)][:5] if isinstance(tags, list) else []
                        matched_customers.append({
                            "id": row[0],
                            "name": row[1],
                            "avatar": row[3] or "",
                            "gender": row[4] or 0,
                            "add_time": row[5] or 0,
                            "owner_userid": row[6] or "",
                            "owner_name": row[7] or "-",
                            "remark": row[8] or "",
                            "tags": customer_tags
                        })
                    except:
                        continue
        
        elif filter_type == "gender":
            # 按性别筛选
            gender_map = {"未知": 0, "男": 1, "女": 2}
            gender_value = gender_map.get(filter_value)
            
            if gender_value is not None:
                cursor.execute("""
                    SELECT id, name, enterprise_tags, avatar, gender, add_time, owner_userid, owner_name, remark
                    FROM customers
                    WHERE gender = ?
                """, (gender_value,))
                
                for row in cursor.fetchall():
                    try:
                        tags = json.loads(row[2]) if row[2] else []
                        customer_tags = [t.get('tag_name', '') for t in tags if isinstance(t, dict)][:5] if isinstance(tags, list) else []
                        matched_customers.append({
                            "id": row[0],
                            "name": row[1],
                            "avatar": row[3] or "",
                            "gender": row[4] or 0,
                            "add_time": row[5] or 0,
                            "owner_userid": row[6] or "",
                            "owner_name": row[7] or "-",
                            "remark": row[8] or "",
                            "tags": customer_tags
                        })
                    except:
                        continue
        
        # 分页
        total = len(matched_customers)
        start = (page - 1) * limit
        end = start + limit
        page_data = matched_customers[start:end]
        
        conn.close()
        
        return {
            "success": True,
            "data": {
                "customers": page_data,
                "total": total,
                "page": page,
                "limit": limit,
                "total_pages": (total + limit - 1) // limit
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/monthly-growth")
async def get_monthly_growth(
    tag_type: str = Query("all", description="标签类型：all/user/agent/partner/supplier/peer/old-agent"),
    months: int = Query(12, ge=1, le=24, description="统计月份数（默认12个月）"),
    token: str = Depends(check_token)
):
    """获取月度客户增长趋势"""
    import json
    from datetime import datetime, timedelta
    from collections import defaultdict
    
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        # 获取当前时间和起始时间
        now = datetime.now()
        start_date = now - timedelta(days=months * 30)
        start_timestamp = int(start_date.timestamp())
        
        # 查询时间范围内的客户
        cursor.execute("""
            SELECT add_time, enterprise_tags
            FROM customers
            WHERE add_time >= ?
            ORDER BY add_time
        """, (start_timestamp,))
        
        # 按月统计
        monthly_data = defaultdict(int)
        
        # 标签关键词映射
        tag_keywords = {
            "user": ["用户", "客户"],
            "agent": ["代理商", "代理"],
            "partner": ["合伙人"],
            "supplier": ["供应商"],
            "peer": ["同行", "上游同行"],
            "old-agent": ["原有老代理", "历史-代理"]
        }
        
        for row in cursor.fetchall():
            add_time = row[0]
            enterprise_tags = row[1]
            
            if not add_time:
                continue
            
            # 判断是否符合标签筛选
            if tag_type != "all":
                keywords = tag_keywords.get(tag_type, [])
                if keywords:
                    try:
                        tags = json.loads(enterprise_tags) if enterprise_tags else []
                        if isinstance(tags, list):
                            has_match = False
                            for tag in tags:
                                if isinstance(tag, dict):
                                    tag_name = tag.get('tag_name', '')
                                    if any(keyword in tag_name for keyword in keywords):
                                        has_match = True
                                        break
                            if not has_match:
                                continue
                    except:
                        continue
            
            # 统计到对应月份
            dt = datetime.fromtimestamp(add_time)
            month_key = dt.strftime('%Y-%m')
            monthly_data[month_key] += 1
        
        conn.close()
        
        # 生成完整的月份列表
        result = []
        for i in range(months):
            month_date = now - timedelta(days=(months - i - 1) * 30)
            month_key = month_date.strftime('%Y-%m')
            count = monthly_data.get(month_key, 0)
            result.append({
                "month": month_key,
                "count": count,
                "display": month_date.strftime('%Y年%m月')
            })
        
        # 计算统计数据
        total_growth = sum(item['count'] for item in result)
        avg_growth = round(total_growth / months, 1) if months > 0 else 0
        max_month = max(result, key=lambda x: x['count']) if result else None
        min_month = min(result, key=lambda x: x['count']) if result else None
        
        return {
            "success": True,
            "data": {
                "months": result,
                "total_growth": total_growth,
                "avg_growth": avg_growth,
                "max_month": max_month,
                "min_month": min_month
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/monthly-growth-by-year")
async def get_monthly_growth_by_year(
    tag_type: str = Query("all", description="标签类型：all/user/agent/partner/supplier/peer/old-agent"),
    year: int = Query(..., description="年份，如2024"),
    token: str = Depends(check_token)
):
    """获取指定年份的月度客户增长趋势（1-12月）"""
    import json
    from datetime import datetime
    from collections import defaultdict
    
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        # 获取该年份的起止时间戳
        start_date = datetime(year, 1, 1)
        end_date = datetime(year, 12, 31, 23, 59, 59)
        start_timestamp = int(start_date.timestamp())
        end_timestamp = int(end_date.timestamp())
        
        # 查询时间范围内的客户
        cursor.execute("""
            SELECT add_time, enterprise_tags
            FROM customers
            WHERE add_time >= ? AND add_time <= ?
            ORDER BY add_time
        """, (start_timestamp, end_timestamp))
        
        # 按月统计
        monthly_data = defaultdict(int)
        
        # 标签关键词映射
        tag_keywords = {
            "user": ["用户", "客户"],
            "agent": ["代理商", "代理"],
            "partner": ["合伙人"],
            "supplier": ["供应商"],
            "peer": ["同行", "上游同行"],
            "old-agent": ["原有老代理", "历史-代理"]
        }
        
        for row in cursor.fetchall():
            add_time = row[0]
            enterprise_tags = row[1]
            
            if not add_time:
                continue
            
            # 判断是否符合标签筛选
            if tag_type != "all":
                keywords = tag_keywords.get(tag_type, [])
                if keywords:
                    try:
                        tags = json.loads(enterprise_tags) if enterprise_tags else []
                        if isinstance(tags, list):
                            has_match = False
                            for tag in tags:
                                if isinstance(tag, dict):
                                    tag_name = tag.get('tag_name', '')
                                    if any(keyword in tag_name for keyword in keywords):
                                        has_match = True
                                        break
                            if not has_match:
                                continue
                    except:
                        continue
            
            # 统计到对应月份
            dt = datetime.fromtimestamp(add_time)
            month_key = f"{year}-{dt.month:02d}"
            monthly_data[month_key] += 1
        
        conn.close()
        
        # 生成1-12月的完整列表
        result = []
        for month in range(1, 13):
            month_key = f"{year}-{month:02d}"
            count = monthly_data.get(month_key, 0)
            result.append({
                "month": month_key,
                "count": count,
                "display": f"{year}年{month:02d}月"
            })
        
        # 计算统计数据
        total_growth = sum(item['count'] for item in result)
        avg_growth = round(total_growth / 12, 1)
        max_month = max(result, key=lambda x: x['count']) if result else None
        min_month = min(result, key=lambda x: x['count']) if result else None
        
        return {
            "success": True,
            "data": {
                "months": result,
                "total_growth": total_growth,
                "avg_growth": avg_growth,
                "max_month": max_month,
                "min_month": min_month
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/employee-ranking")
async def get_employee_ranking(
    tag_type: str = Query("all", description="标签类型：all/user/agent/partner/supplier/peer/old-agent"),
    limit: int = Query(20, ge=1, le=100, description="返回前N名员工"),
    token: str = Depends(check_token)
):
    """获取员工客户数量排行榜"""
    import json
    from collections import defaultdict
    
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        # 标签关键词映射
        tag_keywords = {
            "user": ["用户", "客户"],
            "agent": ["代理商", "代理"],
            "partner": ["合伙人"],
            "supplier": ["供应商"],
            "peer": ["同行", "上游同行"],
            "old-agent": ["原有老代理", "历史-代理"]
        }
        
        # 查询所有客户
        cursor.execute("""
            SELECT owner_userid, owner_name, enterprise_tags
            FROM customers
            WHERE owner_userid IS NOT NULL AND owner_userid != ''
        """)
        
        # 按员工统计
        employee_stats = defaultdict(lambda: {"userid": "", "name": "", "count": 0})
        
        for row in cursor.fetchall():
            owner_userid = row[0]
            owner_name = row[1] or owner_userid
            enterprise_tags = row[2]
            
            # 判断是否符合标签筛选
            if tag_type != "all":
                keywords = tag_keywords.get(tag_type, [])
                if keywords:
                    try:
                        tags = json.loads(enterprise_tags) if enterprise_tags else []
                        if isinstance(tags, list):
                            has_match = False
                            for tag in tags:
                                if isinstance(tag, dict):
                                    tag_name = tag.get('tag_name', '')
                                    if any(keyword in tag_name for keyword in keywords):
                                        has_match = True
                                        break
                            if not has_match:
                                continue
                    except:
                        continue
            
            # 统计
            employee_stats[owner_userid]["userid"] = owner_userid
            employee_stats[owner_userid]["name"] = owner_name
            employee_stats[owner_userid]["count"] += 1
        
        conn.close()
        
        # 转换为列表并排序
        ranking = sorted(employee_stats.values(), key=lambda x: x['count'], reverse=True)[:limit]
        
        # 计算总数和占比
        total = sum(item['count'] for item in ranking)
        for i, item in enumerate(ranking):
            item['rank'] = i + 1
            item['percentage'] = round(item['count'] / total * 100, 1) if total > 0 else 0
        
        return {
            "success": True,
            "data": {
                "ranking": ranking,
                "total": total
            }
        }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

@app.get("/api/customer-portrait/tag-combinations")
async def get_tag_combinations(
    tags: str = Query("", description="标签名称，逗号分隔，例如：四川,代理商"),
    limit: int = Query(10, ge=1, le=50, description="返回热门组合数量"),
    token: str = Depends(check_token)
):
    """获取标签组合分析数据"""
    import json
    from collections import Counter
    from itertools import combinations
    
    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        # 如果指定了标签，查询同时拥有这些标签的客户
        if tags:
            tag_list = [t.strip() for t in tags.split(',') if t.strip()]
            
            cursor.execute("""
                SELECT id, name, enterprise_tags, avatar, gender, add_time, owner_userid, owner_name, remark
                FROM customers
                WHERE enterprise_tags IS NOT NULL AND enterprise_tags != '' AND enterprise_tags != '[]'
            """)
            
            matched_customers = []
            
            for row in cursor.fetchall():
                try:
                    tags_data = json.loads(row[2]) if row[2] else []
                    if isinstance(tags_data, list):
                        # 获取客户的所有标签名
                        customer_tag_names = set()
                        for tag in tags_data:
                            if isinstance(tag, dict):
                                customer_tag_names.add(tag.get('tag_name', ''))
                        
                        # 检查是否包含所有指定标签
                        if all(tag in customer_tag_names for tag in tag_list):
                            customer_tags = [t.get('tag_name', '') for t in tags_data if isinstance(t, dict)][:5]
                            matched_customers.append({
                                "id": row[0],
                                "name": row[1],
                                "avatar": row[3] or "",
                                "gender": row[4] or 0,
                                "add_time": row[5] or 0,
                                "owner_userid": row[6] or "",
                                "owner_name": row[7] or "-",
                                "remark": row[8] or "",
                                "tags": customer_tags
                            })
                except:
                    continue
            
            conn.close()
            
            return {
                "success": True,
                "data": {
                    "combination_tags": tag_list,
                    "customers": matched_customers,
                    "total": len(matched_customers)
                }
            }
        
        # 如果没有指定标签，返回热门标签组合
        else:
            cursor.execute("""
                SELECT enterprise_tags
                FROM customers
                WHERE enterprise_tags IS NOT NULL AND enterprise_tags != '' AND enterprise_tags != '[]'
            """)
            
            # 统计所有标签组合（2个标签的组合）
            combo_counter = Counter()
            
            for row in cursor.fetchall():
                try:
                    tags_data = json.loads(row[0]) if row[0] else []
                    if isinstance(tags_data, list) and len(tags_data) >= 2:
                        tag_names = [tag.get('tag_name', '') for tag in tags_data if isinstance(tag, dict)]
                        # 生成所有2个标签的组合
                        for combo in combinations(sorted(tag_names), 2):
                            combo_counter[combo] += 1
                except:
                    continue
            
            conn.close()
            
            # 获取最热门的组合
            hot_combinations = []
            for combo, count in combo_counter.most_common(limit):
                hot_combinations.append({
                    "tags": list(combo),
                    "count": count,
                    "display": " + ".join(combo)
                })
            
            return {
                "success": True,
                "data": {
                    "hot_combinations": hot_combinations
                }
            }
        
    except Exception as e:
        return {"success": False, "message": str(e)}

# ========== 启动服务 ==========

if __name__ == "__main__":
    import sys
    import uvicorn
    
    # Windows 系统兼容性修复
    if sys.platform == 'win32':
        import asyncio
        # Python 3.8+ Windows 需要设置事件循环策略
        if sys.version_info >= (3, 8):
            asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    
    print("=" * 60)
    print("🚀 企业微信 CRM 管理系统")
    print("=" * 60)
    print(f"✅ 服务已启动: http://{HOST}:{PORT}")
    print(f"📊 数据库路径: {DB_PATH}")
    print(f"🔑 API Token: {API_TOKEN}")
    
    # 启动定时同步任务
    start_scheduled_sync()
    
    print("\n按 Ctrl+C 停止服务")
    print("=" * 60)
    
    uvicorn.run(app, host=HOST, port=PORT, log_level="info")
